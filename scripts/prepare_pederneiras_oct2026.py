"""Build a checked, idempotent October 2026 import from the original LT-PRO-04.

Read-only against the workbook. Writes a reviewable JSON manifest and SQL file;
the SQL refuses to overwrite any existing Pederneiras operational records.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


AREAS = [
    (129, "Usina de Concreto", "MOI"),
    (130, "Produção", "MOD"),
    (131, "Fab Armação e estribos", "MOI"),
    (132, "Encarregado Produção", "MOI"),
    (133, "Sacador núcleo", "MOI"),
]
ACCESSORY_ROWS = set(range(58, 65))
PRODUCT_ROWS = set(range(10, 122))


def integer(value: object, label: str) -> int:
    if value is None:
        return 0
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value or value < 0:
        raise ValueError(f"Quantidade inválida em {label}: {value!r}")
    return int(value)


def decimal(value: object, label: str) -> Decimal:
    if value is None:
        return Decimal("0")
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Mão de obra inválida em {label}: {value!r}")
    amount = Decimal(str(value))
    if amount < 0 or amount % Decimal("0.25"):
        raise ValueError(f"Mão de obra fora do passo de 0,25 em {label}: {value!r}")
    return amount


def product_code(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip() or None


def build_manifest(source: Path, cutoff: date) -> dict:
    book = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    sheet = book["Prod. P. Compl."]
    rows = list(sheet.iter_rows(min_row=1, max_row=134, max_col=35, values_only=True))
    if rows[5][5] != datetime(2026, 10, 1):
        raise ValueError("A aba de produção não é de outubro de 2026")
    if len(set(AREAS)) != len(AREAS):
        raise ValueError("Áreas duplicadas")
    days = []
    skipped = []
    for col in range(4, 35):
        index = col - 1
        cell_date = rows[8][index]
        if not isinstance(cell_date, datetime) or cell_date.date() != date(2026, 10, col - 3):
            raise ValueError(f"Data de cabeçalho inválida na coluna {col}")
        day = cell_date.date()
        items = []
        posts = 0
        accessories = 0
        for row_number in sorted(PRODUCT_ROWS):
            row = rows[row_number - 1]
            quantity = integer(row[index], f"linha {row_number}, dia {day}")
            if not quantity:
                continue
            code = product_code(row[1])
            if not code:
                raise ValueError(f"Produto com quantidade e sem código na linha {row_number}")
            accessory = row_number in ACCESSORY_ROWS
            items.append({"code": code, "quantity": quantity, "accessory": accessory})
            if accessory:
                accessories += quantity
            else:
                posts += quantity
        footer = integer(rows[121][index], f"rodapé, dia {day}")
        labor = []
        for row_number, name, kind in AREAS:
            amount = decimal(rows[row_number - 1][index], f"{name}, dia {day}")
            labor.append({"area": name, "kind": kind, "men": str(amount)})
        labor_total = sum((Decimal(item["men"]) for item in labor), Decimal("0"))
        footer_labor = decimal(rows[133][index], f"total de homens, dia {day}")
        if labor_total != footer_labor:
            raise ValueError(f"Mão de obra não fecha em {day}: {labor_total} vs {footer_labor}")
        if day > cutoff:
            if footer or posts or accessories or labor_total:
                skipped.append({"day": day.isoformat(), "footer": footer, "posts": posts,
                                "accessories": accessories, "men": str(labor_total)})
            continue
        if footer != posts:
            raise ValueError(f"Total de postes não fecha em {day}: {posts} vs {footer}")
        if posts and not labor_total:
            raise ValueError(f"Produção sem mão de obra em {day}")
        if posts:
            days.append({"day": day.isoformat(), "posts": posts,
                         "accessories": accessories, "men": str(labor_total),
                         "items": items, "labor": labor})
        elif accessories or labor_total:
            raise ValueError(f"Lançamento isolado sem postes em {day}")
    codes = [item["code"] for day in days for item in day["items"]]
    if len(codes) != len(set((day["day"], item["code"]) for day in days for item in day["items"])):
        raise ValueError("Código duplicado no mesmo dia")
    return {
        "source_name": source.name,
        "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "cutoff": cutoff.isoformat(),
        "areas": [{"name": name, "kind": kind, "order": position}
                  for position, (_, name, kind) in enumerate(AREAS, 1)],
        "days": days,
        "skipped_future": skipped,
        "totals": {"days": len(days), "posts": sum(day["posts"] for day in days),
                   "accessories": sum(day["accessories"] for day in days),
                   "men": str(sum((Decimal(day["men"]) for day in days), Decimal("0"))),
                   "item_rows": sum(len(day["items"]) for day in days)},
    }


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def import_sql(manifest: dict) -> str:
    areas = manifest["areas"]
    days = manifest["days"]
    area_values = ",\n    ".join(
        f"({sql_literal(area['name'])},{sql_literal(area['kind'])},{area['order']})" for area in areas)
    day_values = ",\n    ".join(
        f"({sql_literal(day['day'])}::date,{day['posts']},{day['accessories']},{day['men']})" for day in days)
    item_values = ",\n    ".join(
        f"({sql_literal(day['day'])}::date,{sql_literal(item['code'])},{item['quantity']})"
        for day in days for item in day["items"])
    labor_values = ",\n    ".join(
        f"({sql_literal(day['day'])}::date,{sql_literal(item['area'])},{item['men']})"
        for day in days for item in day["labor"] if Decimal(item["men"]) > 0)
    source_hash = sql_literal(manifest["source_sha256"])
    expected_posts = manifest["totals"]["posts"]
    expected_accessories = manifest["totals"]["accessories"]
    return f"""\\set ON_ERROR_STOP on
BEGIN;
SET LOCAL lock_timeout = '10s';
SELECT pg_advisory_xact_lock(20261024, 2);
CREATE TEMP TABLE _pdr_areas(name text, kind text, position integer) ON COMMIT DROP;
INSERT INTO _pdr_areas VALUES
    {area_values};
CREATE TEMP TABLE _pdr_days(day date, posts integer, accessories integer, men numeric) ON COMMIT DROP;
INSERT INTO _pdr_days VALUES
    {day_values};
CREATE TEMP TABLE _pdr_items(day date, code text, quantity integer) ON COMMIT DROP;
INSERT INTO _pdr_items VALUES
    {item_values};
CREATE TEMP TABLE _pdr_labor(day date, area text, men numeric) ON COMMIT DROP;
INSERT INTO _pdr_labor VALUES
    {labor_values};
DO $guard$
BEGIN
  IF EXISTS (SELECT 1 FROM public.atividades_mao_de_obra WHERE unidade_id=2)
     OR EXISTS (SELECT 1 FROM public.maodeobra WHERE unidade_id=2)
     OR EXISTS (SELECT 1 FROM public.maodeobra_dias WHERE unidade_id=2)
     OR EXISTS (SELECT 1 FROM producao.apontamento_dia WHERE unidade_id=2) THEN
    RAISE EXCEPTION 'Pederneiras already contains areas or historical records; import aborted';
  END IF;
  IF EXISTS (
    SELECT 1 FROM _pdr_items i LEFT JOIN producao.produto p
      ON p.unidade_id=2 AND p.cod_pro=i.code AND p.ativo
    WHERE p.id IS NULL
  ) THEN RAISE EXCEPTION 'An active catalog product is missing; import aborted'; END IF;
  IF EXISTS (SELECT 1 FROM _pdr_days d LEFT JOIN _pdr_items i USING(day)
    LEFT JOIN producao.produto p ON p.unidade_id=2 AND p.cod_pro=i.code
    GROUP BY d.day,d.posts,d.accessories
    HAVING COALESCE(sum(i.quantity) FILTER (WHERE p.familia='Acessórios'),0)<>d.accessories
      OR COALESCE(sum(i.quantity) FILTER (WHERE p.familia<>'Acessórios'),0)<>d.posts)
  THEN RAISE EXCEPTION 'Catalog families do not match the workbook totals; import aborted'; END IF;
  IF EXISTS (SELECT 1 FROM _pdr_days d LEFT JOIN _pdr_labor l USING(day)
    GROUP BY d.day,d.men HAVING COALESCE(sum(l.men),0)<>d.men)
  THEN RAISE EXCEPTION 'Labor totals do not match the workbook; import aborted'; END IF;
END $guard$;
INSERT INTO public.atividades_mao_de_obra(unidade_id,descricao,tipo,grupo_setor,ordem,ativo)
SELECT 2,name,kind,'02',position,true FROM _pdr_areas ORDER BY position;
INSERT INTO producao.auditoria(unidade_id,entidade,entidade_id,campo,valor_de,valor_para,usuario,origem)
SELECT 2,'area_mao_obra',a.id,'criacao',NULL,a.descricao,'sistema','importacao'
FROM public.atividades_mao_de_obra a WHERE a.unidade_id=2;
INSERT INTO producao.apontamento_dia(unidade_id,data,status,enviado_em,enviado_por,confirmado_em,confirmado_por,origem)
SELECT 2,day,'confirmado',now(),'sistema',now(),'sistema','importacao' FROM _pdr_days ORDER BY day;
INSERT INTO producao.apontamento_item(apontamento_dia_id,produto_id,quantidade)
SELECT d.id,p.id,i.quantity FROM _pdr_items i
JOIN producao.apontamento_dia d ON d.unidade_id=2 AND d.data=i.day
JOIN producao.produto p ON p.unidade_id=2 AND p.cod_pro=i.code;
INSERT INTO producao.auditoria(unidade_id,entidade,entidade_id,campo,valor_de,valor_para,usuario,origem)
SELECT 2,'apontamento_item',i.id,'quantidade','0',i.quantidade::text,'sistema','importacao'
FROM producao.apontamento_item i JOIN producao.apontamento_dia d ON d.id=i.apontamento_dia_id
WHERE d.unidade_id=2;
INSERT INTO public.maodeobra_dias(data,grupo_setor,unidade_id,sem_producao,salvo_em,confirmado_em)
SELECT day,'02',2,false,now(),now() FROM _pdr_days ORDER BY day;
INSERT INTO public.maodeobra(data,atividade_id,unidade_id,quantidade_pessoas)
SELECT l.day,a.id,2,l.men FROM _pdr_labor l
JOIN public.atividades_mao_de_obra a ON a.unidade_id=2 AND a.descricao=l.area;
INSERT INTO producao.auditoria(unidade_id,entidade,entidade_id,campo,valor_de,valor_para,usuario,origem)
SELECT 2,'maodeobra',m.id,'quantidade_pessoas','0',m.quantidade_pessoas::text,'sistema','importacao'
FROM public.maodeobra m WHERE m.unidade_id=2;
INSERT INTO producao.auditoria(unidade_id,entidade,entidade_id,campo,valor_de,valor_para,usuario,origem)
VALUES (2,'importacao_historica',2,'sha256_xlsx',NULL,{source_hash},'sistema','importacao');
DO $verify$
DECLARE actual_posts integer; actual_accessories integer;
BEGIN
  SELECT COALESCE(sum(i.quantidade) FILTER (WHERE p.familia<>'Acessórios'),0),
         COALESCE(sum(i.quantidade) FILTER (WHERE p.familia='Acessórios'),0)
  INTO actual_posts,actual_accessories
  FROM producao.apontamento_item i JOIN producao.apontamento_dia d ON d.id=i.apontamento_dia_id
  JOIN producao.produto p ON p.id=i.produto_id WHERE d.unidade_id=2;
  IF actual_posts<>{expected_posts} OR actual_accessories<>{expected_accessories} THEN
    RAISE EXCEPTION 'Post import totals disagree: posts %, accessories %',actual_posts,actual_accessories;
  END IF;
END $verify$;
COMMIT;
SELECT 'days' AS measure,count(*)::text AS value FROM producao.apontamento_dia WHERE unidade_id=2
UNION ALL SELECT 'areas',count(*)::text FROM public.atividades_mao_de_obra WHERE unidade_id=2
UNION ALL SELECT 'labor_records',count(*)::text FROM public.maodeobra WHERE unidade_id=2
UNION ALL SELECT 'item_records',count(*)::text FROM producao.apontamento_item i
 JOIN producao.apontamento_dia d ON d.id=i.apontamento_dia_id WHERE d.unidade_id=2;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("--cutoff", type=date.fromisoformat, default=date.today())
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    manifest = build_manifest(args.source, args.cutoff)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "pederneiras_oct2026_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (args.output_dir / "pederneiras_oct2026_import.sql").write_text(
        import_sql(manifest), encoding="utf-8")
    print(json.dumps({"source_sha256": manifest["source_sha256"],
                      "totals": manifest["totals"], "skipped_future": manifest["skipped_future"]},
                     ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
