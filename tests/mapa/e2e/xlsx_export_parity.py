"""Compare Mapa XLSX export with the original page on the same test snapshot."""

from __future__ import annotations

import hashlib
import json
import argparse
from pathlib import Path
from urllib.parse import parse_qs, urlsplit
import uuid

from openpyxl import load_workbook
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError, sync_playwright

from operacao_smoke import CHROME, OUT, URL, credentials, login_operacao


SOURCE = Path(__file__).resolve().parents[3] / ".work-mapa-unificado" / "operacao" / "index-before-desktop.html"
DEST = OUT / "xlsx-export"


def workbook_data(path: Path) -> dict:
    book = load_workbook(path, read_only=True, data_only=False)
    result = {}
    for sheet in book:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        if sheet.title == "Resumo":
            for row in rows:
                if row and row[0] == "Gerado em":
                    row[1] = "<hora de geração>"
        result[sheet.title] = rows
    book.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--role", choices=("GERENCIA", "GESTOR"), default="GERENCIA")
    args = parser.parse_args()
    DEST.mkdir(parents=True, exist_ok=True)
    source = SOURCE.read_text(encoding="utf-8")
    account = credentials()
    external = []
    writes = []
    summaries = {}
    selected_day = None
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True, executable_path=CHROME)
        context = browser.new_context(viewport={"width": 1366, "height": 768},
                                      accept_downloads=True, service_workers="block")
        def restrict(route):
            request = route.request
            host = urlsplit(request.url).hostname
            if host != "concrefer.concretrack.com.br":
                external.append(host)
                route.abort()
            elif request.method not in {"GET", "HEAD"}:
                writes.append({"method": request.method, "path": urlsplit(request.url).path})
                if "/unified/login" in request.url:
                    route.continue_()
                else:
                    route.abort()
            else:
                route.continue_()
        context.route("**/*", restrict)
        for variant in ("original", "copia"):
            page = context.new_page()
            page_errors = []
            console_errors = []
            responses = []
            failed_responses = []
            page.on('pageerror', lambda error: page_errors.append(str(error)[:180]))
            page.on('console', lambda message: console_errors.append(message.text[:200])
                    if message.type == 'error' else None)
            page.on('response', lambda response: responses.append({
                    'path': urlsplit(response.url).path, 'status': response.status,
                    'query_excerpt': urlsplit(response.url).query[:250] if response.status >= 400 else ''})
                    if '/rest/v1/' in response.url else None)
            page.on('response', lambda response: failed_responses.append(response)
                    if '/rest/v1/' in response.url and response.status >= 400 else None)
            if variant == "original":
                login_operacao(page, account["MAPA_TEST_USER"],
                               account["MAPA_TEST_PASSWORD"], args.role)
                page.route(URL, lambda route: route.fulfill(status=200,
                    content_type="text/html; charset=utf-8", body=source))
                page.goto(URL, wait_until="domcontentloaded")
            else:
                page.goto(URL, wait_until="domcontentloaded")
            page.wait_for_function("role => state.authUser?.role === role", arg=args.role,
                                   timeout=15000)
            if selected_day is None:
                selected_day = page.evaluate("""async () => {
                  const response=await fetch('/unificado/mapa/api/rest/v1/montagem_poste?select=data_fabricacao,etapa&data_fabricacao=not.is.null&order=data_fabricacao.desc&limit=1000');
                  if (!response.ok) throw new Error('Latest montage date HTTP '+response.status);
                  return (await response.json()).find(row=>!isHistoricoReprovacao(row))?.data_fabricacao;
                }""")
                if not selected_day:
                    raise RuntimeError('No dated montage row in test snapshot')
            page.evaluate("""day => {
              setMode('MONTAGEM_INDICADORES');
              document.getElementById('miDataInicio').value=day;
              document.getElementById('miDataFim').value=day;
            }""", selected_day)
            try:
                with page.expect_download(timeout=25000) as event:
                    page.locator("#miBtnExportarXlsx").click()
            except PlaywrightTimeoutError:
                failed = failed_responses[-1] if failed_responses else None
                raw_filter = parse_qs(urlsplit(failed.url).query).get('id', [''])[0] if failed else ''
                ids = raw_filter.removeprefix('in.(').removesuffix(')').split(',') if raw_filter else []
                invalid = []
                for item in ids:
                    try: uuid.UUID(item)
                    except ValueError: invalid.append(item[:50])
                state = page.evaluate("""() => ({
                  xlsx:!!window.XLSX?.utils,
                  button:document.querySelector('#miBtnExportarXlsx')?.textContent,
                  msg:[...document.querySelectorAll('#msgbox-msg,[role=alert],.msgbox,.msgbox-content')]
                    .map(node=>node.textContent.trim().slice(0,180)).filter(Boolean).slice(-3)
                })""")
                print(json.dumps({"variant": variant, "date": selected_day,
                                  "no_download": state, "page_errors": page_errors,
                                  "console_errors": console_errors[-3:],
                                  "responses": responses[-8:],
                                  "lookup_id_count": len(ids), "invalid_ids": invalid,
                                  "upstream_error": failed.text()[:400] if failed else None}, ensure_ascii=False))
                raise
            download = event.value
            path = DEST / f"montagem-{variant}-{args.role.lower()}.xlsx"
            download.save_as(path)
            data = workbook_data(path)
            encoded = json.dumps(data, sort_keys=True, ensure_ascii=False, default=str).encode()
            summaries[variant] = {"filename": download.suggested_filename,
                                  "sheets": {name: len(rows) for name, rows in data.items()},
                                  "logical_sha256": hashlib.sha256(encoded).hexdigest(),
                                  "data": data}
            page.close()
        context.close()
        browser.close()
    assert not external and not any(item for item in writes if "/unified/login" not in item["path"])
    assert summaries["original"]["filename"] == summaries["copia"]["filename"]
    assert summaries["original"]["data"] == summaries["copia"]["data"]
    for value in summaries.values():
        value.pop("data")
    (DEST / f"report-{args.role.lower()}.json").write_text(
        json.dumps(summaries, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"xlsx_export": "equal", "role": args.role, "date": selected_day,
                      "original": summaries["original"],
                      "copia": summaries["copia"], "external": external,
                      "official_writes": 0}, ensure_ascii=False))


if __name__ == "__main__":
    main()
