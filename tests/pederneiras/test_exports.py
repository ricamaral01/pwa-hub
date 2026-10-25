from datetime import date
from decimal import Decimal
from importlib.util import module_from_spec, spec_from_file_location
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


SOURCE = Path(__file__).resolve().parents[2] / "unificado/pcp/backend/app/services/pederneiras_exports.py"
spec = spec_from_file_location("pederneiras_exports", SOURCE)
exports = module_from_spec(spec)
spec.loader.exec_module(exports)


def test_exports_share_monthly_days_and_totals():
    products = [{"id": 1, "descricao": "Poste A", "cod_pro": "101", "cod_sma": "201"},
                {"id": 2, "descricao": "Poste B", "cod_pro": "102", "cod_sma": "202"}]
    days = [
        {"data": date(2026, 10, 1), "postes": 3, "homens": Decimal("0.75"),
         "mod": Decimal("0.5"), "moi": Decimal("0.25"),
         "postes_homem": Decimal("4"), "volume_m3": None,
         "itens": [{"produto_id": 1, "quantidade": 2}, {"produto_id": 2, "quantidade": 1}]},
        {"data": date(2026, 10, 3), "postes": 1, "homens": Decimal("0.25"),
         "mod": Decimal("0.25"), "moi": Decimal("0"),
         "postes_homem": Decimal("4"), "volume_m3": None,
         "itens": [{"produto_id": 2, "quantidade": 1}]},
    ]
    month = {"unidade": {"nome": "Teste"}, "mes": "2026-10", "dias": days,
             "postes": 4, "homens": Decimal("1")}
    workbook = load_workbook(BytesIO(exports.xlsx(month, products)), data_only=True)
    sheet = workbook.active
    assert [sheet.cell(2, col).value for col in (4, 5)] == ["01/10", "03/10"]
    assert sheet.cell(3, 6).value == 2
    assert sheet.cell(4, 6).value == 2
    assert sheet.cell(5, 6).value == 4
    assert sheet.cell(6, 6).value == 1
    rendered = exports.pdf(month, products)
    assert rendered.startswith(b"%PDF-")
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(rendered)).pages)
    assert "Poste A" in text and "Poste B" in text
    assert "01/10" in text and "03/10" in text
