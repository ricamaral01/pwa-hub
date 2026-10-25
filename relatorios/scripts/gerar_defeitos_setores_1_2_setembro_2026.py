"""Relatório de defeitos dos setores 1 e 2 em setembro de 2026.

Lê montagem_poste e producao no Supabase. A visão principal segue o período
de produção usado pelo dashboard; a segunda segue a data local de inspeção.
"""

from __future__ import annotations

import html
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "pcp_producao_agent"))
import config  # noqa: E402

PERIOD = "2026-09"
BRAZIL = ZoneInfo("America/Sao_Paulo")
SECTORS = ("Setor 1", "Setor 2")
OUT_DIR = ROOT / "relatorios" / "defeitos-postes"
BASE = "relatorio_defeitos_setores_1_2_setembro_2026"
CHECKLIST_LABELS = {
    "homogeneidade_concreto": "Homogeneidade do concreto",
    "falhas_preenchimento": "Falhas de preenchimento",
    "concreto_segregado": "Concreto segregado",
    "grandes_avarias": "Grandes avarias",
    "facao_obstruido": "Facão obstruído",
    "furacao_obstruida": "Furação obstruída (pinos)",
    "armacao_aparente": "Armação aparente",
    "trincas": "Trincas",
    "tubulacao_entupida": "Tubulação entupida",
    "bolhas_excesso": "Bolhas em excesso",
    "bolhas_fora_padrao": "Bolhas fora do padrão",
    "carimbo_identificacao": "Carimbo de identificação",
    "fissuras": "Fissuras",
    "pequenas_avarias": "Pequenas avarias",
    "rebarbas": "Rebarbas",
    "manchas_excessivas": "Manchas excessivas",
    "buchas_fixacao": "Buchas de fixação",
    "prisioneiros_lacre_aterramento": "Prisioneiros (lacre/aterramento)",
    "acabamento_face_exposta": "Acabamento face exposta",
    "acabamento_abas": "Acabamento abas",
    "montagem_do_poste": "Montagem do poste",
    "rebarbas_final": "Rebarbas (inspeção final)",
    "liberacao_qualidade": "Liberação da qualidade",
    "codificacao_poste": "Codificação do poste",
    "aterramento": "Limpeza do aterramento",
    "lacre": "Limpeza do lacre",
}
REASON_LABELS = {
    "A": "Falha na concretagem",
    "B": "Cano entupido",
    "C": "Excesso de bolhas",
    "D": "Problema na caixa do relógio ou disjuntor",
    "E": "Furos obstruídos ou faltando",
    "F": "Identificação do poste",
    "G": "Trincas superficiais",
    "H": "Bolhas nas caixas",
    "I": "Trincas na extensão do poste",
    "J": "Pequenas avarias",
    "K": "Manchas brancas superficiais",
    "L": "Buchas das caixas ou fixação",
    "M": "Parafuso lacre da caixa do medidor",
    "O": "Falha no concreto",
    "P": "Grandes avarias",
    "Q": "Rebarbas",
    "R": "Rugosidade",
    "S": "Acabamento das quinas",
}


def fetch_all(table: str, params: dict) -> list[dict]:
    headers = {"apikey": config.SUPABASE_KEY, "Authorization": f"Bearer {config.SUPABASE_KEY}"}
    result = []
    offset = 0
    while True:
        query = {**params, "order": "id.asc", "limit": "1000", "offset": str(offset)}
        response = requests.get(
            f"{config.SUPABASE_URL}/rest/v1/{table}", headers=headers,
            params=query, timeout=45,
        )
        response.raise_for_status()
        batch = response.json()
        result.extend(batch)
        if len(batch) < 1000:
            return result
        offset += 1000


def parsed_utc(value: str | None) -> datetime | None:
    if not value:
        return None
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def inspection_date(row: dict) -> str:
    checklists = row.get("checklists") or {}
    raw = checklists.get("dia_inspecao") if isinstance(checklists, dict) else None
    if raw:
        try:
            return datetime.strptime(raw.strip(), "%d/%m/%Y").date().isoformat()
        except ValueError:
            pass
    finished = parsed_utc(row.get("finalizado_em"))
    return finished.astimezone(BRAZIL).date().isoformat() if finished else ""


def rejected_items(row: dict) -> list[tuple[str, str, str]]:
    items = []
    checklists = row.get("checklists") or {}
    if not isinstance(checklists, dict):
        return items
    for section, answers in checklists.items():
        if not isinstance(answers, dict):
            continue
        for key, answer in answers.items():
            if isinstance(answer, str) and answer.strip().casefold() == "nao":
                items.append((section, key, CHECKLIST_LABELS.get(key, key.replace("_", " ").capitalize())))
    return items


def defect_items(row: dict) -> list[tuple[str, str, str, str]]:
    negatives = rejected_items(row)
    if negatives:
        return [(section, key, label, "Checklist: resposta não") for section, key, label in negatives]
    status = str(row.get("status_montagem") or "").strip().upper()
    reason = str(row.get("motivo_recusa") or "").strip().upper()
    if reason:
        return [(
            "motivo_recusa", reason,
            REASON_LABELS.get(reason, f"Motivo de recusa: {reason}"),
            "Motivo de recusa; sem item 'não' no checklist",
        )]
    if status in {"R", "RR", "REPROVADO", "RETRABALHO"}:
        return [("status_montagem", status, "Reprovação/retrabalho sem defeito especificado", "Status sem motivo detalhado")]
    return []


def fetch_montage():
    base = {
        "select": "id,data_fabricacao,setor,forma_numero,modelo,status_montagem,motivo_recusa,etapa,finalizado_em,checklists",
        "setor": 'in.("Setor 1","Setor 2")',
    }
    produced = fetch_all("montagem_poste", {
        **base, "data_fabricacao": "gte.2026-09-01",
        "and": "(data_fabricacao.lt.2026-10-01)",
    })
    inspected = fetch_all("montagem_poste", {
        **base, "finalizado_em": "gte.2026-09-01T03:00:00Z",
        "and": "(finalizado_em.lt.2026-10-01T03:00:00Z)",
    })
    # A data local no checklist é a referência final quando preenchida.
    union = {row["id"]: row for row in inspected}
    union.update({row["id"]: row for row in produced})
    return list(union.values()), len(produced), len(inspected)


def fetch_production_for(rows: list[dict]) -> dict:
    dates = sorted({row["data_fabricacao"] for row in rows if row.get("data_fabricacao")})
    if not dates:
        return {}
    production = fetch_all("producao", {
        "select": "id,data_fabricacao,setor,forma,modelo,codigo_poste,descricao_poste,codigo_produto,tipo_concreto,status,data_hora",
        "setor": 'in.("Setor 1","Setor 2")',
        "data_fabricacao": f"in.({','.join(dates)})",
    })
    by_key = defaultdict(list)
    for row in production:
        by_key[(row["data_fabricacao"], row["setor"], str(row["forma"] or "").strip().upper())].append(row)
    for matches in by_key.values():
        matches.sort(
            key=lambda row: (
                row.get("tipo_concreto") not in (None, "", "INSPECIONADO"),
                row.get("status") in ("LIBERADO", "CONCRETADO"),
                row.get("data_hora") or "",
            ), reverse=True,
        )
    return by_key


def build_occurrences(montage: list[dict], production: dict) -> list[dict]:
    occurrences = []
    for row in montage:
        defects = defect_items(row)
        if not defects:
            continue
        key = (row["data_fabricacao"], row["setor"], str(row["forma_numero"] or "").strip().upper())
        produced = (production.get(key) or [None])[0]
        concrete = produced.get("tipo_concreto") if produced else None
        if concrete == "INSPECIONADO":
            concrete = None
        for section, defect_key, label, evidence in defects:
            occurrences.append({
                "sector": row["setor"],
                "model": row.get("modelo") or (produced or {}).get("modelo") or "Não informado",
                "form": row.get("forma_numero") or "",
                "defect": label,
                "defect_key": defect_key,
                "section": section,
                "evidence": evidence,
                "production_date": row.get("data_fabricacao") or "",
                "inspection_date": inspection_date(row),
                "concrete": concrete or "Não informado",
                "status": row.get("status_montagem") or "",
                "reason": row.get("motivo_recusa") or "",
                "montage_id": row["id"],
                "production_id": (produced or {}).get("id") or "",
            })
    return sorted(occurrences, key=lambda r: (
        r["sector"], r["model"], r["production_date"], r["form"], r["defect"],
    ))


HEADERS = [
    "Setor", "Tipo de poste", "Forma", "Defeito", "Data de produção", "Data de inspeção",
    "Tipo de concreto", "Status", "Motivo de recusa", "Origem do defeito", "Seção", "Campo",
    "ID montagem", "ID produção",
]


def row_values(row):
    return [
        row["sector"], row["model"], row["form"], row["defect"],
        row["production_date"], row["inspection_date"], row["concrete"],
        row["status"], row["reason"], row["evidence"], row["section"],
        row["defect_key"], row["montage_id"], row["production_id"],
    ]


def style_sheet(sheet, widths):
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="173A5E")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def write_xlsx(produced_rows, inspected_rows, produced_count, inspected_count, fetched_at, path):
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Resumo"
    overview.append(["Defeitos — setores 1 e 2 — setembro de 2026", "Valor"])
    overview.append(["Data da consulta (horário de Brasília)", fetched_at.strftime("%d/%m/%Y %H:%M")])
    overview.append(["Postes com registro de montagem e produção em setembro", produced_count])
    overview.append(["Postes com defeito produzidos em setembro", len({r["montage_id"] for r in produced_rows})])
    overview.append(["Ocorrências em postes produzidos em setembro", len(produced_rows)])
    overview.append(["Postes com inspeção concluída em setembro", inspected_count])
    overview.append(["Postes com defeito inspecionados em setembro", len({r["montage_id"] for r in inspected_rows})])
    overview.append(["Ocorrências inspecionadas em setembro", len(inspected_rows)])
    overview.append(["Ocorrências produzidas em agosto e inspecionadas em setembro", sum(r["production_date"].startswith("2026-08") for r in inspected_rows)])
    overview.append(["Tipo de concreto", "producao.tipo_concreto, cruzado por data + setor + forma"])
    overview.append(["Fonte dos defeitos", "montagem_poste.checklists; motivo_recusa/status quando não há item rejeitado"])
    overview.append(["Data de inspeção", "checklists.dia_inspecao; caso ausente, finalizado_em em America/Sao_Paulo"])
    overview.append(["Critério principal", "data_fabricacao de 01/09/2026 a 30/09/2026"])
    overview.column_dimensions["A"].width = 67
    overview.column_dimensions["B"].width = 85
    style_sheet(overview, (67, 85))

    grouped = workbook.create_sheet("Por poste e defeito")
    grouped.append(["Setor", "Tipo de poste", "Defeito", "Tipo de concreto", "Ocorrências"])
    counts = Counter((r["sector"], r["model"], r["defect"], r["concrete"]) for r in produced_rows)
    for key, count in sorted(counts.items()):
        grouped.append([*key, count])
    style_sheet(grouped, (16, 25, 36, 28, 16))

    for title, rows in (("Produzidos setembro", produced_rows), ("Inspecionados setembro", inspected_rows)):
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row_values(row))
        style_sheet(sheet, (16, 25, 17, 37, 20, 20, 28, 13, 18, 44, 24, 32, 60, 18))
    workbook.save(path)


def fmt_date(iso):
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%d/%m/%Y") if iso else "Não registrada"


def html_table(rows):
    body = "".join(
        "<tr>" + "".join(f"<td>{html.escape(str(value))}</td>" for value in (
            r["sector"], r["model"], r["form"], r["defect"],
            fmt_date(r["production_date"]), fmt_date(r["inspection_date"]),
            r["concrete"], r["status"], r["evidence"],
        )) + "</tr>" for r in rows
    )
    return (
        "<div class='scroll'><table><thead><tr><th>Setor</th><th>Tipo de poste</th><th>Forma</th>"
        "<th>Defeito</th><th>Produção</th><th>Inspeção</th><th>Concreto</th><th>Status</th>"
        "<th>Origem do defeito</th></tr></thead><tbody>" + body + "</tbody></table></div>"
    )


def write_html(produced_rows, inspected_rows, fetched_at, path):
    sector = Counter(r["sector"] for r in produced_rows)
    defects = Counter((r["sector"], r["defect"]) for r in produced_rows)
    summary = "".join(
        f"<tr><td>{html.escape(s)}</td><td>{html.escape(d)}</td><td class='num'>{n}</td></tr>"
        for (s, d), n in sorted(defects.items())
    )
    content = f"""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1"><title>Defeitos · Setembro de 2026</title>
<style>body{{font:14px/1.45 Arial,sans-serif;color:#1d2939;max-width:1500px;margin:28px auto;padding:0 22px}}h1,h2{{color:#173a5e}}h1{{margin-bottom:4px}}.sub{{color:#667085}}
.cards{{display:flex;gap:12px;flex-wrap:wrap;margin:20px 0}}.card{{background:#edf4fa;border-left:4px solid #23669b;padding:12px 16px;min-width:170px}}.card strong{{display:block;font-size:23px}}
.note{{background:#fff4df;border-left:4px solid #d59a32;padding:12px 16px;margin:18px 0}}table{{border-collapse:collapse;width:100%;font-size:12px;margin:10px 0 25px}}
th,td{{border:1px solid #d4dde5;padding:6px 8px;text-align:left;vertical-align:top}}th{{background:#173a5e;color:#fff;position:sticky;top:0}}
tr:nth-child(even){{background:#f5f8fa}}.num{{text-align:right}}.scroll{{overflow:auto}}@media print{{th{{position:static}}.scroll{{overflow:visible}}tr{{break-inside:avoid}}}}</style></head><body>
<h1>Defeitos dos setores 1 e 2 — Setembro de 2026</h1><p class="sub">Dados operacionais consultados em {fetched_at.strftime('%d/%m/%Y %H:%M')} (Brasília).</p>
<div class="cards"><div class="card"><strong>{len(produced_rows)}</strong>ocorrências em postes produzidos em setembro</div>
<div class="card"><strong>{len({r['montage_id'] for r in produced_rows})}</strong>postes com defeito</div>
<div class="card"><strong>{sector['Setor 1']}</strong>ocorrências no setor 1</div><div class="card"><strong>{sector['Setor 2']}</strong>ocorrências no setor 2</div></div>
<div class="note"><strong>Critério principal:</strong> data de produção entre 01/09/2026 e 30/09/2026. Cada resposta “não” do checklist é uma ocorrência de defeito.
Quando há reprovação ou motivo de recusa sem resposta “não”, o motivo é apresentado como ocorrência identificada pelo status. A data de inspeção vem de <code>dia_inspecao</code> e, se ausente, de <code>finalizado_em</code> no horário de Brasília.
O tipo de concreto vem de <code>producao.tipo_concreto</code>, ligado pela data de produção, setor e forma. O relatório é um retrato da base no horário indicado.</div>
<h2>Resumo por defeito</h2><table><thead><tr><th>Setor</th><th>Defeito</th><th>Ocorrências</th></tr></thead><tbody>{summary}</tbody></table>
<h2>Postes produzidos em setembro</h2>{html_table(produced_rows)}
<h2>Postes inspecionados em setembro</h2><p>Esta visão inclui {sum(r['production_date'].startswith('2026-08') for r in inspected_rows)} ocorrências de postes produzidos em agosto e inspecionados em setembro.</p>
{html_table(inspected_rows)}
<p class="sub">Fonte: tabelas <code>montagem_poste</code> e <code>producao</code> do Supabase do projeto. Os IDs de origem estão no arquivo Excel.</p>
</body></html>"""
    path.write_text(content, encoding="utf-8")


def main():
    fetched_at = datetime.now(timezone.utc).astimezone(BRAZIL)
    montage, produced_count, inspected_count = fetch_montage()
    eligible = [row for row in montage if defect_items(row)]
    production = fetch_production_for(eligible)
    occurrences = build_occurrences(eligible, production)
    produced_rows = [row for row in occurrences if row["production_date"].startswith(PERIOD)]
    inspected_rows = [row for row in occurrences if row["inspection_date"].startswith(PERIOD)]
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx_path = OUT_DIR / f"{BASE}.xlsx"
    html_path = OUT_DIR / f"{BASE}.html"
    write_xlsx(produced_rows, inspected_rows, produced_count, inspected_count, fetched_at, xlsx_path)
    write_html(produced_rows, inspected_rows, fetched_at, html_path)
    check = load_workbook(xlsx_path, read_only=True)
    assert check["Produzidos setembro"].max_row == len(produced_rows) + 1
    assert check["Inspecionados setembro"].max_row == len(inspected_rows) + 1
    print(f"Consulta: {fetched_at.strftime('%d/%m/%Y %H:%M %Z')}")
    print(f"Produção setembro: {len(produced_rows)} ocorrências em {len({r['montage_id'] for r in produced_rows})} postes")
    print(f"Inspeção setembro: {len(inspected_rows)} ocorrências em {len({r['montage_id'] for r in inspected_rows})} postes")
    print(f"Concreto não encontrado: {sum(r['concrete'] == 'Não informado' for r in produced_rows)} ocorrência(s)")
    print(xlsx_path)
    print(html_path)


if __name__ == "__main__":
    main()
