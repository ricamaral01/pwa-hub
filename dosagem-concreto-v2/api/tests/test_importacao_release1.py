from pathlib import Path

import openpyxl

from conftest import rpc


REFERENCE = Path(__file__).resolve().parents[3] / "docs/dosagem-v2/referencia"


def test_rica_simulation_effective_import_and_idempotence(ambiente):
    """R1-IMPORT-RICA: import four real Rica v4 tabs, warning about the physical volume."""
    client, _, _ = ambiente
    book = openpyxl.load_workbook(REFERENCE / "Dosagem_Concreto_Rica_v4.xlsx", read_only=True, data_only=True)
    names = ["R10 (2)", "FILLER V2", "MARCOS", "TESTE AGREGADO V2"]
    records = []
    for name in names:
        sheet = book[name]
        items = []
        for column in ("C", "F", "H", "K", "N", "P", "R", "V", "X", "Y", "Z"):
            mass = sheet[f"{column}15"].value
            density = sheet[f"{column}14"].value
            if isinstance(mass, (int, float)) and mass > 0:
                assert isinstance(density, (int, float)) and density > 0
                items.append({"column": column, "massKgM3": mass, "densityKgDm3": density})
        air = sheet["X9"].value
        volume = sum(item["massKgM3"] / item["densityKgDm3"] for item in items) + air * 1000
        assert volume > 0
        records.append({"tipo": "ensaios", "origem": "rica-v4", "id_antigo": name,
                        "documento": {"id": f"rica-{names.index(name)}", "unidade": "campinas", "nome": name,
                                      "tipo": "conferencia_traco", "status": "conferencia",
                                      "dados": {"itens": items, "arIncorporadoPct": air * 100,
                                                "volumeCalculadoL": volume, "statusEngenharia": "conferencia_sem_aprovacao"}}})
    assert any(abs(record["documento"]["dados"]["volumeCalculadoL"] - 1000) > .01 for record in records)
    request = {"simulacao": True, "registros": records, "unidadePadrao": "campinas"}
    preview = rpc(client, "importacao.legado", request)
    assert preview.status_code == 200, preview.json()
    assert preview.json()["dados"]["gravados"] == 0
    assert preview.json()["dados"]["convertidos"] == 4
    imported = rpc(client, "importacao.legado", {**request, "simulacao": False})
    assert imported.status_code == 200, imported.json()
    assert imported.json()["dados"]["gravados"] == 4
    repeated = rpc(client, "importacao.legado", {**request, "simulacao": False})
    assert repeated.status_code == 200, repeated.json()
    assert repeated.json()["dados"]["gravados"] == 0
    assert repeated.json()["dados"]["duplicados"] == 4


def test_udemy_and_abrams_reference_fixtures(ambiente):
    client, _, _ = ambiente
    udemy = openpyxl.load_workbook(REFERENCE / "1.Dosagem-Udemy.xlsx", read_only=True, data_only=True)["Dosagem "]
    assert (udemy["E9"].value, udemy["I9"].value, udemy["J9"].value) == (30, 189, 321)
    curve_book = openpyxl.load_workbook(REFERENCE / "RECURSOS-CURVADEABRAMS-UDEMY.xlsx", read_only=True, data_only=True)
    measured = curve_book["Curva de abrams 1"]
    points = [{"ac": measured[f"D{row}"].value, "resistencia28MPa": measured[f"F{row}"].value,
               "aguaL": measured[f"E{row}"].value} for row in range(6, 11)]
    assert len(points) == 5 and all(point["ac"] > 0 and point["resistencia28MPa"] > 0 for point in points)
    incomplete = curve_book["Curva de abrams 2"]
    assert all(incomplete[f"F{row}"].value == 0 for row in range(5, 10))
    records = [
        {"tipo": "ensaios", "origem": "udemy-dosagem", "id_antigo": "Dosagem ",
         "documento": {"id": "udemy-ref", "unidade": "campinas", "tipo": "referencia_dosagem",
                       "status": "conferencia", "dados": {"fckMPa": 30, "aguaL": 189, "cimentoKg": 321}}},
        {"tipo": "ensaios", "origem": "curva-abrams-udemy", "id_antigo": "Curva de abrams 1",
         "documento": {"id": "abrams-ref", "unidade": "campinas", "tipo": "curva_abrams_referencia",
                       "status": "conferencia", "dados": {"pontos": points}}},
    ]
    preview = rpc(client, "importacao.legado", {"simulacao": True, "registros": records})
    assert preview.status_code == 200 and preview.json()["dados"]["convertidos"] == 2
    imported = rpc(client, "importacao.legado", {"simulacao": False, "registros": records})
    assert imported.status_code == 200 and imported.json()["dados"]["gravados"] == 2
    repeated = rpc(client, "importacao.legado", {"simulacao": False, "registros": records})
    assert repeated.status_code == 200 and repeated.json()["dados"]["duplicados"] == 2
