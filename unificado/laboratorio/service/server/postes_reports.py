"""Print and spreadsheet reports for FL-PRO-22."""
from __future__ import annotations
from html import escape
from io import BytesIO
import sys

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .postes_db import Equipment, Responsible


def fmt(value, digits=1):
    return "—" if value is None else f"{float(value):,.{digits}f}".replace(",", "_").replace(".", ",").replace("_", ".")


def date_br(value):
    return '—' if value is None else value.strftime('%d/%m/%Y')


def meta(record, session):
    equipment = session.get(Equipment, record.equipamento_id) if record.equipamento_id else None
    inspector = session.get(Responsible, record.inspetor_id) if record.inspetor_id else None
    engineer = session.get(Responsible, record.engenheiro_id) if record.engenheiro_id else None
    return equipment, inspector, engineer


def report_html(record, session):
    equipment, inspector, engineer = meta(record, session)
    pole = record.poste_snapshot
    par = record.parametros_calculados
    result = record.resultados
    steps = result.get("etapas", {})
    labels = ["Engaste e acomodação", "Momento fletor MA", "Elasticidade · nominal", "Elasticidade · excepcional", "Ruptura"]
    refs = [f'{fmt(par.get("carga_nom_daN"), 0)} daN · 1 min', "NBR 8451-2 · 3 min",
            f'{fmt(par.get("flecha_max_mm"))} mm · 3 min', f'{fmt(par.get("residual_max_mm"))} mm · 3 min',
            f'{fmt(par.get("ruptura_min_daN"), 0)} daN']
    readings = ["Confirmada" if record.leituras.get("e1") else "Pendente",
                "N/A" if steps.get("e2", {}).get("status") == "N/A" else f'{fmt(steps.get("e2", {}).get("ma_daNm"),0)} daN·m · fiss. {fmt(steps.get("e2", {}).get("fissura_mm"),2)} mm · {record.leituras.get("e2", {}).get("ok", "—")}',
                f'{fmt(steps.get("e3", {}).get("valor_mm"))} mm · fiss. {fmt(steps.get("e3", {}).get("fissura_mm"), 2)} mm',
                f'{fmt(steps.get("e4", {}).get("valor_mm"))} mm · fiss. {fmt(steps.get("e4", {}).get("fissura_mm"), 2)} mm',
                f'{fmt(steps.get("e5", {}).get("valor_daN"), 0)} daN']
    percentages = ["—", "—", fmt(steps.get("e3", {}).get("uso_pct"), 0) + "% do limite",
                   fmt(steps.get("e4", {}).get("uso_pct"), 0) + "% do limite",
                   fmt(steps.get("e5", {}).get("pct_min"), 0) + "% do mínimo"]
    rows = "".join(f'<tr><td>{i}</td><td>{escape(labels[i-1])}</td><td>{escape(refs[i-1])}</td>'
                   f'<td>{escape(readings[i-1])}</td><td>{escape(percentages[i-1])}</td>'
                   f'<td>{escape(steps.get(f"e{i}", {}).get("status", "—"))}</td></tr>' for i in range(1, 6))
    absorb = result.get("absorcao", {})
    absorb_rows = "".join(f'<tr><td>CP {i}</td><td>{fmt(v, 2)}%</td></tr>' for i, v in enumerate(absorb.get("valores_pct", []), 1))
    def esc(value): return escape(str(value if value is not None else "—"))
    return f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>FL-PRO-22 · Amostra {esc(record.numero_amostra)}</title>
<style>@page{{size:A4 landscape;margin:12mm}}*{{box-sizing:border-box}}body{{font:12px Arial,sans-serif;color:#182c35;margin:22px}}
h1{{font-size:18px;margin:0}}h2{{font-size:13px;margin:18px 0 7px}}header{{display:flex;justify-content:space-between;border-bottom:3px solid #114c5a;padding-bottom:10px}}
.sub{{color:#536a70}}.grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:7px;margin:14px 0}}.cell{{border:1px solid #c6d2d2;padding:7px}}
.cell b{{display:block;font-size:10px;color:#50666b;margin-bottom:4px}}table{{width:100%;border-collapse:collapse;margin:8px 0}}td,th{{border:1px solid #bac8c9;padding:6px;text-align:left}}th{{background:#e8efee}}
.verdict{{font-weight:bold;font-size:16px;border:2px solid #114c5a;padding:10px;margin-top:13px}}footer{{margin-top:18px;border-top:1px solid #a8b9b9;padding-top:9px;display:grid;grid-template-columns:1fr 1fr 1fr;gap:15px}}
.print{{float:right;border:0;background:#114c5a;color:white;padding:9px 14px;cursor:pointer}}@media print{{.print{{display:none}}body{{margin:0}}}}</style></head><body>
<button class="print" onclick="window.print()">Imprimir</button><header><div><h1>CONCREFER · ENSAIO EM BANCADA DE POSTES</h1><div class="sub">FL-PRO-22 · Rev. 00 · NBR 8451-1/2/3/4/5</div></div><div>Amostra <b>{esc(record.numero_amostra)}</b><br>Revisão {record.revisao}</div></header>
<div class="grid"><div class="cell"><b>POSTE</b>{esc(pole.get('tipo'))} {fmt(pole.get('altura_m'))}/{fmt(pole.get('resistencia_daN'),0)} · face {esc(pole.get('face'))}</div>
<div class="cell"><b>DATA DO ENSAIO</b>{date_br(record.data_ensaio)}</div><div class="cell"><b>FABRICAÇÃO</b>{date_br(record.data_fabricacao)}</div><div class="cell"><b>FÁBRICA</b>{esc(record.fabrica)}</div>
<div class="cell"><b>ENGASTE</b>{fmt(par.get('engaste_cm'),0)} cm</div><div class="cell"><b>CARRO DO TOPO</b>{fmt(par.get('carro_cm'),0)} cm</div>
<div class="cell"><b>CARGA NOMINAL</b>{fmt(par.get('carga_nom_daN'),0)} daN</div><div class="cell"><b>CARGA EXCEPCIONAL</b>{fmt(par.get('carga_exc_daN'),0)} daN</div></div>
<h2>Sequência de ensaio</h2><table><thead><tr><th>Etapa</th><th>Procedimento</th><th>Referência/limite</th><th>Leitura</th><th>%</th><th>Status</th></tr></thead><tbody>{rows}</tbody></table>
<div class="verdict">Resultado da amostra: {esc(record.veredito)} · Inspeção geral: {'OK' if record.inspecao_geral_ok else 'NÃO OK'} · Critério v{record.criterio_versao}</div>
<h2>Absorção de água · NBR 8451-4</h2><table><thead><tr><th>Corpo de prova</th><th>Absorção</th></tr></thead><tbody>{absorb_rows}<tr><td>Média</td><td>{fmt(absorb.get('media_pct'),2)}% · {esc(absorb.get('status'))}</td></tr></tbody></table>
<footer><div><b>Dinamômetro</b><br>{esc(equipment.descricao if equipment else None)} · {fmt(equipment.capacidade_kgf,0) if equipment else '—'} kgf<br>Série {esc(equipment.numero_serie if equipment else None)} · Cert. {esc(equipment.certificado if equipment else None)}<br>Calibração: {date_br(equipment.data_calibracao if equipment else None)}</div>
<div><b>Inspetor</b><br>{esc(inspector.nome if inspector else None)}<br><br>Assinatura: ____________________</div><div><b>Engenheiro responsável</b><br>{esc(engineer.nome if engineer else None)}<br>CREA {esc(engineer.registro if engineer else None)}<br>Assinatura: ____________________</div></footer>
<p class="sub">Registro {esc(record.id)} · {esc(record.criado_em)} · por {esc(record.criado_por)}. Critérios e parâmetros preservados no banco.</p></body></html>'''


def report_pdf(record, session):
    if sys.platform == 'win32':
        return _reportlab_pdf(record, session)
    try:
        from weasyprint import HTML
        return HTML(string=report_html(record, session)).write_pdf()
    except (ImportError, OSError):
        # On Windows WeasyPrint may be installed without native Pango/GObject.
        return _reportlab_pdf(record, session)


def _reportlab_pdf(record, session):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from pathlib import Path

    font_file = next((path for path in (
        Path('C:/Windows/Fonts/arial.ttf'),
        Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
    ) if path.exists()), None)
    if font_file and 'ArialLocal' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('ArialLocal', str(font_file)))
    font = 'ArialLocal' if font_file else 'Helvetica'
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                                 topMargin=25, bottomMargin=25)
    styles = getSampleStyleSheet()
    heading = ParagraphStyle('heading_local', parent=styles['Heading1'], fontName=font, fontSize=14, textColor=colors.HexColor('#114C5A'))
    normal = ParagraphStyle('normal_local', parent=styles['Normal'], fontName=font, fontSize=8, leading=11)
    title = ParagraphStyle('title_local', parent=normal, alignment=TA_CENTER, fontSize=10)
    equipment, inspector, engineer = meta(record, session)
    pole, par, steps = record.poste_snapshot, record.parametros_calculados, record.resultados['etapas']
    story = [Paragraph('CONCREFER · ENSAIO EM BANCADA DE POSTES', heading),
             Paragraph(f'FL-PRO-22 · Rev. 00 · Amostra {escape(record.numero_amostra)} · Revisão {record.revisao}', normal), Spacer(1, 10)]
    info = [[f"Poste: {pole.get('tipo')} {fmt(pole.get('altura_m'))}/{fmt(pole.get('resistencia_daN'),0)} · face {pole.get('face')}",
             f"Ensaio: {date_br(record.data_ensaio)}", f"Engaste: {fmt(par.get('engaste_cm'),0)} cm",
             f"Carro: {fmt(par.get('carro_cm'),0)} cm"]]
    info_table = Table(info, colWidths=[240, 130, 150, 150])
    info_table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor('#E8EFEE')),
                                    ('BOX',(0,0),(-1,-1),.5,colors.HexColor('#BAC8C9')),
                                    ('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),8),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]))
    story += [info_table, Spacer(1, 12)]
    labels = ['Engaste e acomodação','Momento fletor MA','Elasticidade nominal','Elasticidade excepcional','Ruptura']
    references = [f"{fmt(par['carga_nom_daN'],0)} daN · 1 min", 'NBR 8451-2 · 3 min',
                  f"{fmt(par['flecha_max_mm'])} mm · 3 min", f"{fmt(par['residual_max_mm'])} mm · 3 min",
                  f"{fmt(par['ruptura_min_daN'],0)} daN"]
    table_data = [['Etapa','Procedimento','Referência/limite','Leitura','%','Status']]
    for i in range(1,6):
        step = steps[f'e{i}']
        if i==1: reading='Confirmada' if record.leituras.get('e1') else 'Pendente'
        elif i==2: reading='N/A' if step['status']=='N/A' else f"{fmt(step.get('ma_daNm'),0)} daN·m · fiss. {fmt(step.get('fissura_mm'),2)} mm"
        elif i in (3,4): reading=f"{fmt(step.get('valor_mm'))} mm · fiss. {fmt(step.get('fissura_mm'),2)} mm"
        else: reading=f"{fmt(step.get('valor_daN'),0)} daN"
        pct = f"{fmt(step.get('uso_pct'),0)}%" if i in (3,4) else f"{fmt(step.get('pct_min'),0)}%" if i==5 else '—'
        table_data.append([str(i),labels[i-1],references[i-1],reading,pct,step['status']])
    table = Table(table_data, colWidths=[38,178,145,175,70,78], repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#114C5A')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,-1),font),
        ('FONTSIZE',(0,0),(-1,-1),8),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#BAC8C9')),
        ('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]))
    absorb = record.resultados.get('absorcao', {})
    story += [table, Spacer(1, 13), Paragraph(f'Resultado da amostra: {record.veredito} · Critério v{record.criterio_versao}', title),
              Paragraph(f"Fábrica: {escape(record.fabrica)} · Fabricação: {date_br(record.data_fabricacao)} · Inspeção geral: {'OK' if record.inspecao_geral_ok else 'NÃO OK'}", normal),
              Paragraph(f"Absorção de água: média {fmt(absorb.get('media_pct'),2)}% · {absorb.get('status','PENDENTE')}", normal), Spacer(1, 15)]
    footer = [[f"Dinamômetro: {equipment.descricao if equipment else '—'}\nSérie: {equipment.numero_serie if equipment else '—'}\nCertificado: {equipment.certificado if equipment else '—'}",
               f"Inspetor: {inspector.nome if inspector else '—'}\nAssinatura: ____________________",
               f"Engenheiro: {engineer.nome if engineer else '—'}\nCREA: {engineer.registro if engineer else '—'}\nAssinatura: ____________________"]]
    footer_table = Table(footer, colWidths=[250,220,230])
    footer_table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,-1),font),('FONTSIZE',(0,0),(-1,-1),8),
        ('BOX',(0,0),(-1,-1),.4,colors.HexColor('#BAC8C9')),
        ('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    story.append(footer_table)
    document.build(story)
    return output.getvalue()


def report_xlsx(records, session):
    wb = Workbook()
    ws = wb.active
    ws.title = "FL-PRO-22"
    ws.merge_cells("A1:Y1")
    ws["A1"] = "CONCREFER · FL-PRO-22 · ENSAIO EM BANCADA DE POSTES"
    ws.merge_cells("A2:Y2")
    ws["A2"] = "Rev. 00 · NBR 8451-1/2/3/4/5 · resultados calculados com os critérios salvos em cada ensaio"
    headers = ["Amostra", "Tipo", "Face", "Altura (m)", "Resist. (daN)", "Fabricação", "Ensaio",
        "Carro (cm)", "Engaste (cm)", "Carga nom. (daN)", "Acomodação", "Momento MA", "Status MA",
        "Fissura nom. (mm)", "Flecha max. (mm)", "Flecha lida (mm)", "% L nom.", "Carga exc. (daN)",
        "Residual max. (mm)", "Residual lido (mm)", "% L res.", "Fissura res. (mm)",
        "Ruptura mín. (daN)", "Ruptura lida (daN)", "Resultado"]
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value)
        cell.fill = PatternFill("solid", fgColor="114C5A")
        cell.font = Font(color="FFFFFF", bold=True, size=9)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = 14 if col > 3 else 11
    ws.row_dimensions[4].height = 34
    for index, r in enumerate(records, 5):
        p, par, steps = r.poste_snapshot, r.parametros_calculados, r.resultados.get("etapas", {})
        e3, e4, e5 = steps.get("e3", {}), steps.get("e4", {}), steps.get("e5", {})
        values = [r.numero_amostra, p.get("tipo"), p.get("face"), p.get("altura_m"), p.get("resistencia_daN"),
            date_br(r.data_fabricacao), date_br(r.data_ensaio), par.get("carro_cm"), par.get("engaste_cm"),
            par.get("carga_nom_daN"), steps.get("e1", {}).get("status"),
            r.leituras.get("e2", {}).get("ma_daNm"), steps.get("e2", {}).get("status"),
            e3.get("fissura_mm"), par.get("flecha_max_mm"), e3.get("valor_mm"), e3.get("pct_L"),
            par.get("carga_exc_daN"), par.get("residual_max_mm"), e4.get("valor_mm"), e4.get("pct_L"),
            e4.get("fissura_mm"), par.get("ruptura_min_daN"), e5.get("valor_daN"), r.veredito]
        for col, value in enumerate(values, 1):
            cell = ws.cell(index, col, value)
            cell.border = Border(bottom=Side(style="hair", color="CAD5D5"))
            if col in (17, 21): cell.number_format = "0.00"
        if index % 2:
            for cell in ws[index]: cell.fill = PatternFill("solid", fgColor="EDF4F2")
    foot = max(7, len(records) + 7)
    ws.merge_cells(start_row=foot, start_column=1, end_row=foot, end_column=25)
    ws.cell(foot, 1, "FL-PRO-22 · Relatório de ensaios finalizados · Dinamômetro, certificado, inspetor e engenheiro constam no laudo individual.")
    ws.freeze_panes = "H5"
    ws.auto_filter.ref = f"A4:Y{max(5, len(records)+4)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "1:4"
    details = wb.create_sheet('Equipamento e responsáveis')
    detail_headers = ['Amostra', 'Revisão', 'Fábrica', 'Inspetor', 'Engenheiro', 'CREA',
                      'Dinamômetro', 'Capacidade (kgf)', 'Série', 'Certificado', 'Calibração', 'Critério v']
    details.append(detail_headers)
    for cell in details[1]:
        cell.fill = PatternFill('solid', fgColor='114C5A')
        cell.font = Font(color='FFFFFF', bold=True)
    for r in records:
        eq, insp, eng = meta(r, session)
        details.append([r.numero_amostra, r.revisao, r.fabrica,
            insp.nome if insp else '', eng.nome if eng else '', eng.registro if eng else '',
            eq.descricao if eq else '', float(eq.capacidade_kgf) if eq else None,
            eq.numero_serie if eq else '', eq.certificado if eq else '',
            date_br(eq.data_calibracao) if eq else '', r.criterio_versao])
    for col in range(1, 13): details.column_dimensions[get_column_letter(col)].width = 24
    details.freeze_panes = 'A2'
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
