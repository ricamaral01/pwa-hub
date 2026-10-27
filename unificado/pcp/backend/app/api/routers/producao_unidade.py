"""Unit-scoped production API. Existing Campinas routes are untouched."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Literal
from zoneinfo import ZoneInfo
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.core.security import get_current_user
from app.models.usuario import NivelAcesso, Usuario
from app.services.pederneiras_catalog import parse_catalog, preview
from app.services.pederneiras_rules import (
    NON_POST_FAMILY, assert_can_confirm, counts_as_post, monthly_target, posts_per_person, production_summary,
    validate_day, validate_people, validate_quantity,
)

router = APIRouter(prefix="/api/producao/{unidade}", tags=["producao-unidade"])


class ItemInput(BaseModel):
    produto_id: int = Field(gt=0)
    quantidade: int = Field(ge=0)


class ApontamentoInput(BaseModel):
    itens: list[ItemInput]
    acao: Literal["rascunho", "enviar"] = "rascunho"
    origem: Literal["desktop", "pwa"] = "desktop"


class AreaInput(BaseModel):
    atividade_id: int = Field(gt=0)
    homens: Decimal = Field(ge=0)
    observacao: str | None = Field(default=None, max_length=2000)


class MaoDeObraInput(BaseModel):
    areas: list[AreaInput]
    origem: Literal["desktop", "pwa"] = "desktop"


class ProductInput(BaseModel):
    familia: str = Field(min_length=1, max_length=120)
    descricao: str = Field(min_length=1, max_length=255)
    letra: str | None = Field(default=None, max_length=20)
    cod_pro: str | None = Field(default=None, max_length=50)
    cod_sma: str | None = Field(default=None, max_length=50)
    volume_m3: Decimal | None = Field(default=None, ge=0)
    pendencia: str | None = None
    ativo: bool = True


class ReopenInput(BaseModel):
    motivo: str = Field(min_length=5, max_length=500)
    origem: Literal["desktop", "pwa"] = "desktop"


class AreaInputCadastro(BaseModel):
    descricao: str = Field(min_length=1, max_length=255)
    tipo: Literal["MOD", "MOI"]
    ordem: int = Field(ge=0, le=9999)
    ativo: bool = True


def _unit(db: Session, slug: str) -> dict:
    row = db.execute(text("SELECT id,codigo,nome,passo_homens FROM producao.unidade WHERE codigo=:c AND ativo"), {"c": slug}).mappings().first()
    if not row:
        raise HTTPException(404, "Unidade não encontrada")
    return dict(row)


def _authorize(db: Session, unit: dict, user: Usuario, *, write: bool = False, supervisor: bool = False) -> str:
    # Core administrators remain masters of every production unit.
    if user.username.casefold() in {"ricardo", "marcelo", "alex", "alex51", "philippe"} and user.nivel_acesso == NivelAcesso.ADMIN:
        return "supervisor"
    row = db.execute(text("""
        SELECT papel FROM producao.usuario_unidade
        WHERE usuario_id=:user_id AND unidade_id=:unit_id AND ativo
    """), {"user_id": user.id, "unit_id": unit["id"]}).first()
    role = row[0] if row else None
    if role is None or (write and role not in {"gestor", "supervisor", "operador"}) or (supervisor and role != "supervisor"):
        raise HTTPException(403, "Sem acesso a esta unidade ou ação")
    return role


def _day(day: date) -> date:
    try:
        return validate_day(day)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc


def _labor_group(unit: dict) -> str:
    # Two-character code fits the legacy day key; unit 1 keeps existing 12/34.
    if unit["id"] > 99 or unit["id"] in {12, 34}:
        raise HTTPException(409, "Código de unidade excede a chave de mão de obra")
    return f"{unit['id']:02d}"


def _audit(db: Session, unit: dict, entity: str, entity_id: int, field: str,
           before: object, after: object, user: Usuario, origin: str):
    if before == after:
        return
    db.execute(text("""
      INSERT INTO producao.auditoria
        (unidade_id,entidade,entidade_id,campo,valor_de,valor_para,usuario,origem)
      VALUES (:uid,:entity,:eid,:field,:before,:after,:user,:origin)
    """), {"uid": unit["id"], "entity": entity, "eid": entity_id,
            "field": field, "before": None if before is None else str(before),
            "after": None if after is None else str(after), "user": user.username,
            "origin": origin})


def _product_rows(db: Session, unit: dict):
    return [dict(row) for row in db.execute(text("""
      SELECT id,familia,letra,descricao,cod_pro,cod_sma,volume_m3,ativo,pendencia
      FROM producao.produto WHERE unidade_id=:uid
      ORDER BY (pendencia IS NULL),familia,descricao,id
    """), {"uid": unit["id"]}).mappings()]


@router.get("/produtos")
def products(unidade: str, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user)
    return {"unidade": unit, "produtos": _product_rows(db, unit)}


@router.get("/produtos/modelo.xlsx")
def product_template(unidade: str, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Produtos"
    sheet.append(["Família", "Letra", "Descrição", "Cód. Pro", "Cód. SMA",
                  "m³/peça", "Ativo", "Situação"])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17324F")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)
    for column, width in {"A": 25, "B": 10, "C": 55, "D": 16, "E": 16,
                          "F": 14, "G": 12, "H": 45}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="modelo_produtos_{unit["codigo"]}.xlsx"'})


@router.post("/produtos/importar")
async def import_products(unidade: str, arquivo: UploadFile = File(...), aplicar: bool = Query(False),
                          db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=aplicar)
    try:
        rows = parse_catalog(arquivo.filename or "", await arquivo.read())
    except (ValueError, UnicodeError) as exc:
        raise HTTPException(422, str(exc)) from exc
    differences = preview(rows, _product_rows(db, unit))
    counts = {action: sum(item["acao"] == action for item in differences)
              for action in ("criar", "atualizar", "igual", "conflito")}
    if not aplicar:
        return {"aplicado": False, "resumo": counts, "diferencas": differences}
    if counts["conflito"]:
        raise HTTPException(409, "Resolva os conflitos na prévia antes de aplicar")
    try:
        for difference in differences:
            action = difference["acao"]
            if action == "igual":
                continue
            values = difference["produto"]
            if action == "criar":
                product_id = db.execute(text("""
                  INSERT INTO producao.produto
                    (unidade_id,familia,letra,descricao,cod_pro,cod_sma,volume_m3,ativo,pendencia)
                  VALUES (:uid,:familia,:letra,:descricao,:cod_pro,:cod_sma,:volume_m3,:ativo,:pendencia)
                  RETURNING id
                """), {"uid": unit["id"], **values}).scalar_one()
                _audit(db, unit, "produto", product_id, "importacao", None,
                       values["descricao"], user, "desktop")
            else:
                product_id = difference["produto_id"]
                db.execute(text("""
                  UPDATE producao.produto SET familia=:familia,letra=:letra,descricao=:descricao,
                    cod_pro=:cod_pro,cod_sma=:cod_sma,volume_m3=:volume_m3,ativo=:ativo,
                    pendencia=:pendencia,atualizado_em=now()
                  WHERE id=:id AND unidade_id=:uid
                """), {"id": product_id, "uid": unit["id"], **values})
                for field, change in difference["alteracoes"].items():
                    _audit(db, unit, "produto", product_id, field,
                           change["de"], change["para"], user, "desktop")
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"aplicado": True, "resumo": counts, "diferencas": differences}


def _product_values(payload: ProductInput) -> dict:
    values = payload.model_dump()
    for key in ("familia", "descricao", "letra", "cod_pro", "cod_sma", "pendencia"):
        if isinstance(values[key], str):
            values[key] = values[key].strip() or None
    if not values["familia"] or not values["descricao"]:
        raise HTTPException(422, "Família e descrição obrigatórias")
    if values["volume_m3"] is not None and values["volume_m3"] > Decimal("9999.9999"):
        raise HTTPException(422, "Volume por peça excede o limite")
    return values


@router.post("/produtos", status_code=201)
def create_product(unidade: str, payload: ProductInput,
                   db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    values = _product_values(payload)
    if values["cod_pro"] and db.execute(text("""
      SELECT 1 FROM producao.produto WHERE unidade_id=:uid AND cod_pro=:code
    """), {"uid": unit["id"], "code": values["cod_pro"]}).first():
        raise HTTPException(409, "Código Pro já cadastrado na unidade")
    try:
        row = db.execute(text("""
          INSERT INTO producao.produto
            (unidade_id,familia,descricao,letra,cod_pro,cod_sma,volume_m3,pendencia,ativo)
          VALUES (:uid,:familia,:descricao,:letra,:cod_pro,:cod_sma,:volume_m3,:pendencia,:ativo)
          RETURNING id
        """), {"uid": unit["id"], **values}).scalar_one()
        _audit(db, unit, "produto", row, "criacao", None, values["descricao"], user, "desktop")
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"id": row, **values}


@router.patch("/produtos/{produto_id}")
def update_product(unidade: str, produto_id: int, payload: ProductInput,
                   db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    values = _product_values(payload)
    try:
        old = db.execute(text("""
          SELECT * FROM producao.produto WHERE id=:id AND unidade_id=:uid FOR UPDATE
        """), {"id": produto_id, "uid": unit["id"]}).mappings().first()
        if not old:
            raise HTTPException(404, "Produto não encontrado na unidade")
        if values["cod_pro"] and db.execute(text("""
          SELECT 1 FROM producao.produto WHERE unidade_id=:uid AND cod_pro=:code AND id<>:id
        """), {"uid": unit["id"], "code": values["cod_pro"], "id": produto_id}).first():
            raise HTTPException(409, "Código Pro já cadastrado na unidade")
        db.execute(text("""
          UPDATE producao.produto SET familia=:familia,descricao=:descricao,letra=:letra,
            cod_pro=:cod_pro,cod_sma=:cod_sma,volume_m3=:volume_m3,
            pendencia=:pendencia,ativo=:ativo,atualizado_em=now()
          WHERE id=:id AND unidade_id=:uid
        """), {"id": produto_id, "uid": unit["id"], **values})
        for key, value in values.items():
            _audit(db, unit, "produto", produto_id, key, old[key], value, user, "desktop")
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"id": produto_id, **values}


def _production(db: Session, unit: dict, day: date) -> dict:
    marker = db.execute(text("""
      SELECT id,status,enviado_em,enviado_por,confirmado_em,confirmado_por
      FROM producao.apontamento_dia WHERE unidade_id=:uid AND data=:day
    """), {"uid": unit["id"], "day": day}).mappings().first()
    rows = [dict(row) for row in db.execute(text("""
      SELECT i.id,i.produto_id,i.quantidade,p.descricao,p.familia,p.volume_m3
      FROM producao.apontamento_item i
      JOIN producao.apontamento_dia d ON d.id=i.apontamento_dia_id
      JOIN producao.produto p ON p.id=i.produto_id AND p.unidade_id=d.unidade_id
      WHERE d.unidade_id=:uid AND d.data=:day
      ORDER BY p.familia,p.descricao
    """), {"uid": unit["id"], "day": day}).mappings()]
    summary = production_summary([(r["quantidade"], r["volume_m3"]) for r in rows
                                  if counts_as_post(r["familia"])])
    return {"unidade": unit, "data": day, "status": marker["status"] if marker else "rascunho",
            "enviado_em": marker["enviado_em"] if marker else None,
            "confirmado_em": marker["confirmado_em"] if marker else None,
            "liberacao_prevista": date.fromordinal(day.toordinal() + 3),
            "itens": rows, **summary}


@router.get("/apontamentos/{dia}")
def get_day(unidade: str, dia: date, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user)
    return _production(db, unit, _day(dia))


@router.get("/apontamentos/{dia}/referencia")
def production_reference(unidade: str, dia: date, db: Session = Depends(get_db),
                         user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user)
    _day(dia)
    ids = [row[0] for row in db.execute(text("""
      SELECT d.id FROM producao.apontamento_dia d
      JOIN producao.apontamento_item i ON i.apontamento_dia_id=d.id
      JOIN producao.produto p ON p.id=i.produto_id
      WHERE d.unidade_id=:uid AND d.data<:day
        AND lower(trim(p.familia))<>lower(:non_post)
      GROUP BY d.id,d.data HAVING sum(i.quantidade)>0
      ORDER BY d.data DESC LIMIT 15
    """), {"uid": unit["id"], "day": dia, "non_post": NON_POST_FAMILY})]
    if not ids:
        return {"dias": 0, "medias": {}}
    rows = db.execute(text("""
      SELECT i.produto_id,sum(i.quantidade) AS total
      FROM producao.apontamento_item i WHERE i.apontamento_dia_id = ANY(:ids)
      GROUP BY i.produto_id
    """), {"ids": ids}).mappings()
    return {"dias": len(ids), "medias": {str(row["produto_id"]):
            round(row["total"] / len(ids), 2) for row in rows}}


@router.put("/apontamentos/{dia}")
def save_day(unidade: str, dia: date, payload: ApontamentoInput,
             db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    _day(dia)
    ids = [item.produto_id for item in payload.itens]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, "Produto repetido no apontamento")
    for item in payload.itens:
        validate_quantity(item.quantidade)
        if not db.execute(text("SELECT 1 FROM producao.produto WHERE id=:id AND unidade_id=:uid AND ativo"),
                          {"id": item.produto_id, "uid": unit["id"]}).first():
            raise HTTPException(422, f"Produto {item.produto_id} não pertence à unidade ou está inativo")
    try:
        db.execute(text("""
          INSERT INTO producao.apontamento_dia(unidade_id,data) VALUES (:uid,:day)
          ON CONFLICT (unidade_id,data) DO NOTHING
        """), {"uid": unit["id"], "day": dia})
        marker = db.execute(text("""
          SELECT id,status FROM producao.apontamento_dia
          WHERE unidade_id=:uid AND data=:day FOR UPDATE
        """), {"uid": unit["id"], "day": dia}).mappings().one()
        if marker["status"] == "confirmado":
            raise HTTPException(409, "Dia confirmado; reabertura exige supervisor e motivo")
        existing = {r["produto_id"]: dict(r) for r in db.execute(text("""
          SELECT id,produto_id,quantidade FROM producao.apontamento_item
          WHERE apontamento_dia_id=:day_id FOR UPDATE
        """), {"day_id": marker["id"]}).mappings()}
        incoming = {item.produto_id: item.quantidade for item in payload.itens}
        for product_id in set(existing) | set(incoming):
            old = existing.get(product_id)
            new_qty = incoming.get(product_id, 0)
            if old:
                if old["quantidade"] != new_qty:
                    db.execute(text("UPDATE producao.apontamento_item SET quantidade=:qty WHERE id=:id"),
                               {"qty": new_qty, "id": old["id"]})
                    _audit(db, unit, "apontamento_item", old["id"], "quantidade",
                           old["quantidade"], new_qty, user, payload.origem)
            elif new_qty:
                item_id = db.execute(text("""
                  INSERT INTO producao.apontamento_item(apontamento_dia_id,produto_id,quantidade)
                  VALUES (:day_id,:pid,:qty) RETURNING id
                """), {"day_id": marker["id"], "pid": product_id, "qty": new_qty}).scalar_one()
                _audit(db, unit, "apontamento_item", item_id, "quantidade", 0, new_qty, user, payload.origem)
        db.execute(text("""
          UPDATE producao.apontamento_dia SET status=:status,origem=:origin,
            enviado_em=CASE WHEN :status='enviado' THEN now() ELSE enviado_em END,
            enviado_por=CASE WHEN :status='enviado' THEN :username ELSE enviado_por END
          WHERE id=:id
        """), {"status": "enviado" if payload.acao == "enviar" else "rascunho",
                "origin": payload.origem, "username": user.username, "id": marker["id"]})
        db.commit()
    except Exception:
        db.rollback()
        raise
    return _production(db, unit, dia)


@router.post("/apontamentos/{dia}/repetir-ontem")
def repeat_last_day(unidade: str, dia: date,
                    origem: Literal["desktop", "pwa"] = Query("desktop"),
                    db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    _day(dia)
    if db.execute(text("""
      SELECT 1 FROM producao.apontamento_dia WHERE unidade_id=:uid AND data=:day
    """), {"uid": unit["id"], "day": dia}).first():
        raise HTTPException(409, "Dia já iniciado; repetição não sobrescreve lançamentos")
    source = db.execute(text("""
      SELECT d.id,d.data FROM producao.apontamento_dia d
      JOIN producao.apontamento_item i ON i.apontamento_dia_id=d.id
      JOIN producao.produto p ON p.id=i.produto_id
      WHERE d.unidade_id=:uid AND d.data<:day
        AND lower(trim(p.familia))<>lower(:non_post)
      GROUP BY d.id,d.data HAVING sum(i.quantidade)>0
      ORDER BY d.data DESC LIMIT 1
    """), {"uid": unit["id"], "day": dia, "non_post": NON_POST_FAMILY}).mappings().first()
    if not source:
        raise HTTPException(404, "Não há dia anterior com produção")
    items = [ItemInput(produto_id=row["produto_id"], quantidade=row["quantidade"])
             for row in db.execute(text("""
               SELECT i.produto_id,i.quantidade FROM producao.apontamento_item i
               JOIN producao.produto p ON p.id=i.produto_id
               WHERE i.apontamento_dia_id=:id AND i.quantidade>0 AND p.ativo
             """), {"id": source["id"]}).mappings()]
    result = save_day(unidade, dia, ApontamentoInput(itens=items, origem=origem), db, user)
    result["copiado_de"] = source["data"]
    return result


@router.post("/apontamentos/{dia}/reabrir")
def reopen_day(unidade: str, dia: date, payload: ReopenInput,
               db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, supervisor=True)
    _day(dia)
    try:
        row = db.execute(text("""
          SELECT id,status FROM producao.apontamento_dia
          WHERE unidade_id=:uid AND data=:day FOR UPDATE
        """), {"uid": unit["id"], "day": dia}).mappings().first()
        if not row or row["status"] != "confirmado":
            raise HTTPException(409, "Somente dia confirmado pode ser reaberto")
        db.execute(text("""
          UPDATE producao.apontamento_dia SET status='rascunho',confirmado_em=NULL,confirmado_por=NULL
          WHERE id=:id
        """), {"id": row["id"]})
        _audit(db, unit, "apontamento_dia", row["id"], "reabertura",
               "confirmado", f"rascunho: {payload.motivo.strip()}", user, payload.origem)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return _production(db, unit, dia)


def _labor(db: Session, unit: dict, day: date) -> dict:
    group = _labor_group(unit)
    marker = db.execute(text("""
      SELECT salvo_em,confirmado_em FROM public.maodeobra_dias
      WHERE unidade_id=:uid AND data=:day AND grupo_setor=:grp
    """), {"uid": unit["id"], "day": day, "grp": group}).mappings().first()
    rows = [dict(row) for row in db.execute(text("""
      SELECT a.id AS atividade_id,a.descricao,a.tipo,a.ordem,a.ativo,
             m.id AS registro_id,COALESCE(m.quantidade_pessoas,0) AS homens,m.observacao
      FROM public.atividades_mao_de_obra a
      LEFT JOIN public.maodeobra m ON m.atividade_id=a.id AND m.data=:day AND m.unidade_id=:uid
      WHERE a.unidade_id=:uid ORDER BY a.ordem,a.descricao
    """), {"uid": unit["id"], "day": day}).mappings()]
    total = sum((Decimal(str(row["homens"])) for row in rows), Decimal("0"))
    mod = sum((Decimal(str(row["homens"])) for row in rows if row["tipo"] == "MOD"), Decimal("0"))
    return {"data": day, "areas": rows, "homens": total, "mod": mod, "moi": total - mod,
            "salvo_em": marker["salvo_em"] if marker else None,
            "confirmado_em": marker["confirmado_em"] if marker else None}


@router.get("/areas")
def get_areas(unidade: str, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user)
    return {"unidade": unit, "areas": [dict(row) for row in db.execute(text("""
      SELECT id,descricao,tipo,ordem,ativo FROM public.atividades_mao_de_obra
      WHERE unidade_id=:uid ORDER BY ordem,descricao,id
    """), {"uid": unit["id"]}).mappings()]}


@router.post("/areas", status_code=201)
def create_area(unidade: str, payload: AreaInputCadastro,
                db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, supervisor=True)
    name = payload.descricao.strip()
    if db.execute(text("""
      SELECT 1 FROM public.atividades_mao_de_obra
      WHERE unidade_id=:uid AND lower(descricao)=lower(:name)
    """), {"uid": unit["id"], "name": name}).first():
        raise HTTPException(409, "Área já cadastrada na unidade")
    try:
        area_id = db.execute(text("""
          INSERT INTO public.atividades_mao_de_obra
            (unidade_id,descricao,tipo,grupo_setor,ordem,ativo)
          VALUES (:uid,:name,:kind,:grp,:sort,:active) RETURNING id
        """), {"uid": unit["id"], "name": name, "kind": payload.tipo,
               "grp": _labor_group(unit), "sort": payload.ordem, "active": payload.ativo}).scalar_one()
        _audit(db, unit, "area_mao_obra", area_id, "criacao", None, name, user, "desktop")
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"id": area_id, **payload.model_dump(), "descricao": name}


@router.patch("/areas/{area_id}")
def update_area(unidade: str, area_id: int, payload: AreaInputCadastro,
                db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, supervisor=True)
    name = payload.descricao.strip()
    try:
        old = db.execute(text("""
          SELECT id,descricao,tipo,ordem,ativo FROM public.atividades_mao_de_obra
          WHERE id=:id AND unidade_id=:uid FOR UPDATE
        """), {"id": area_id, "uid": unit["id"]}).mappings().first()
        if not old:
            raise HTTPException(404, "Área não encontrada na unidade")
        if db.execute(text("""
          SELECT 1 FROM public.atividades_mao_de_obra
          WHERE unidade_id=:uid AND lower(descricao)=lower(:name) AND id<>:id
        """), {"uid": unit["id"], "name": name, "id": area_id}).first():
            raise HTTPException(409, "Área já cadastrada na unidade")
        values = {**payload.model_dump(), "descricao": name}
        db.execute(text("""
          UPDATE public.atividades_mao_de_obra SET descricao=:descricao,tipo=:tipo,
            ordem=:ordem,ativo=:ativo WHERE id=:id AND unidade_id=:uid
        """), {"id": area_id, "uid": unit["id"], **values})
        for key, value in values.items():
            _audit(db, unit, "area_mao_obra", area_id, key, old[key], value, user, "desktop")
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"id": area_id, **values}


@router.get("/mao-de-obra/mensal")
def monthly_labor(unidade: str, mes: str = Query(pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
                  db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    data = monthly(unidade, mes, db, user)
    return {"unidade": data["unidade"], "mes": mes,
            "dias": [{key: day[key] for key in ("data", "postes", "homens", "mod", "moi", "postes_homem", "areas")}
                     for day in data["dias"]]}


@router.get("/mao-de-obra/{dia}")
def get_labor(unidade: str, dia: date, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user)
    return _labor(db, unit, _day(dia))


@router.put("/mao-de-obra/{dia}")
def save_labor(unidade: str, dia: date, payload: MaoDeObraInput,
               db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    _day(dia)
    ids = [area.atividade_id for area in payload.areas]
    if len(ids) != len(set(ids)):
        raise HTTPException(422, "Área repetida")
    for area in payload.areas:
        try:
            validate_people(area.homens, step=unit["passo_homens"])
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc
        if not db.execute(text("SELECT 1 FROM public.atividades_mao_de_obra WHERE id=:id AND unidade_id=:uid AND ativo"),
                          {"id": area.atividade_id, "uid": unit["id"]}).first():
            raise HTTPException(422, f"Área {area.atividade_id} inválida para a unidade")
    group = _labor_group(unit)
    try:
        db.execute(text("""
          INSERT INTO public.maodeobra_dias(data,grupo_setor,unidade_id,sem_producao)
          VALUES (:day,:grp,:uid,false)
          ON CONFLICT (data,grupo_setor) DO NOTHING
        """), {"day": dia, "grp": group, "uid": unit["id"]})
        marker = db.execute(text("""
          SELECT id,confirmado_em FROM public.maodeobra_dias
          WHERE data=:day AND grupo_setor=:grp AND unidade_id=:uid FOR UPDATE
        """), {"day": dia, "grp": group, "uid": unit["id"]}).mappings().one()
        if marker["confirmado_em"]:
            raise HTTPException(409, "Mão de obra confirmada; reabertura exige supervisor")
        existing = {r["atividade_id"]: dict(r) for r in db.execute(text("""
          SELECT m.id,m.atividade_id,m.quantidade_pessoas,m.observacao
          FROM public.maodeobra m JOIN public.atividades_mao_de_obra a ON a.id=m.atividade_id
          WHERE m.data=:day AND m.unidade_id=:uid AND a.unidade_id=:uid FOR UPDATE OF m
        """), {"day": dia, "uid": unit["id"]}).mappings()}
        incoming = {area.atividade_id: area for area in payload.areas}
        for activity_id in set(existing) | set(incoming):
            old = existing.get(activity_id)
            new = incoming.get(activity_id)
            count = new.homens if new else Decimal("0")
            note = new.observacao if new else None
            if old:
                if old["quantidade_pessoas"] != count or old["observacao"] != note:
                    db.execute(text("""
                      UPDATE public.maodeobra SET quantidade_pessoas=:count,observacao=:note WHERE id=:id
                    """), {"count": count, "note": note, "id": old["id"]})
                    _audit(db, unit, "maodeobra", old["id"], "quantidade_pessoas",
                           old["quantidade_pessoas"], count, user, payload.origem)
                    _audit(db, unit, "maodeobra", old["id"], "observacao",
                           old["observacao"], note, user, payload.origem)
            elif count or note:
                record_id = db.execute(text("""
                  INSERT INTO public.maodeobra(data,atividade_id,unidade_id,quantidade_pessoas,observacao)
                  VALUES (:day,:aid,:uid,:count,:note) RETURNING id
                """), {"day": dia, "aid": activity_id, "uid": unit["id"],
                       "count": count, "note": note}).scalar_one()
                _audit(db, unit, "maodeobra", record_id, "quantidade_pessoas", 0, count, user, payload.origem)
        db.execute(text("""
          UPDATE public.maodeobra_dias SET salvo_em=now(),salvo_por_usuario_id=:user_id
          WHERE id=:id
        """), {"user_id": user.id, "id": marker["id"]})
        db.commit()
    except Exception:
        db.rollback()
        raise
    return _labor(db, unit, dia)


@router.post("/mao-de-obra/{dia}/repetir-ontem")
def repeat_last_labor(unidade: str, dia: date, origem: Literal["desktop", "pwa"] = Query("desktop"),
                      db: Session = Depends(get_db),
                      user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    _day(dia)
    if db.execute(text("""
      SELECT 1 FROM public.maodeobra_dias
      WHERE unidade_id=:uid AND data=:day AND grupo_setor=:grp
    """), {"uid": unit["id"], "day": dia, "grp": _labor_group(unit)}).first():
        raise HTTPException(409, "Dia já iniciado; repetição não sobrescreve lançamentos")
    source = db.execute(text("""
      SELECT m.data FROM public.maodeobra m
      JOIN public.atividades_mao_de_obra a ON a.id=m.atividade_id AND a.unidade_id=m.unidade_id
      WHERE m.unidade_id=:uid AND m.data<:day
      GROUP BY m.data HAVING sum(m.quantidade_pessoas)>0
      ORDER BY m.data DESC LIMIT 1
    """), {"uid": unit["id"], "day": dia}).scalar()
    if not source:
        raise HTTPException(404, "Não há dia anterior com mão de obra")
    areas = [AreaInput(atividade_id=row["atividade_id"], homens=row["quantidade_pessoas"],
                       observacao=row["observacao"])
             for row in db.execute(text("""
               SELECT m.atividade_id,m.quantidade_pessoas,m.observacao
               FROM public.maodeobra m JOIN public.atividades_mao_de_obra a ON a.id=m.atividade_id
               WHERE m.unidade_id=:uid AND m.data=:day AND a.unidade_id=:uid AND a.ativo
             """), {"uid": unit["id"], "day": source}).mappings()]
    result = save_labor(unidade, dia, MaoDeObraInput(areas=areas, origem=origem), db, user)
    result["copiado_de"] = source
    return result


@router.post("/mao-de-obra/{dia}/reabrir")
def reopen_labor(unidade: str, dia: date, payload: ReopenInput,
                 db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, supervisor=True)
    _day(dia)
    try:
        marker = db.execute(text("""
          SELECT id,confirmado_em FROM public.maodeobra_dias
          WHERE unidade_id=:uid AND data=:day AND grupo_setor=:grp FOR UPDATE
        """), {"uid": unit["id"], "day": dia, "grp": _labor_group(unit)}).mappings().first()
        if not marker or not marker["confirmado_em"]:
            raise HTTPException(409, "Somente mão de obra confirmada pode ser reaberta")
        db.execute(text("""
          UPDATE public.maodeobra_dias SET confirmado_em=NULL,confirmado_por_usuario_id=NULL
          WHERE id=:id
        """), {"id": marker["id"]})
        _audit(db, unit, "maodeobra_dia", marker["id"], "reabertura",
               marker["confirmado_em"], payload.motivo.strip(), user, payload.origem)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return _labor(db, unit, dia)


@router.post("/mao-de-obra/{dia}/confirmar")
def confirm_labor(unidade: str, dia: date, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    _day(dia)
    try:
        group = _labor_group(unit)
        marker = db.execute(text("""
          SELECT id FROM public.maodeobra_dias
          WHERE unidade_id=:uid AND data=:day AND grupo_setor=:grp FOR UPDATE
        """), {"uid": unit["id"], "day": dia, "grp": group}).first()
        if not marker:
            raise HTTPException(409, "Salve a mão de obra antes de confirmar")
        db.execute(text("""
          UPDATE public.maodeobra_dias SET confirmado_em=now(),confirmado_por_usuario_id=:user_id
          WHERE id=:id AND confirmado_em IS NULL
        """), {"user_id": user.id, "id": marker[0]})
        db.commit()
    except Exception:
        db.rollback()
        raise
    return _labor(db, unit, dia)


@router.post("/apontamentos/{dia}/confirmar")
def confirm_day(unidade: str, dia: date, db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user, write=True)
    _day(dia)
    try:
        marker = db.execute(text("""
          SELECT id,status FROM producao.apontamento_dia
          WHERE unidade_id=:uid AND data=:day FOR UPDATE
        """), {"uid": unit["id"], "day": dia}).mappings().first()
        if not marker or marker["status"] != "enviado":
            raise HTTPException(409, "Envie o apontamento antes de confirmar")
        posts = db.execute(text("""
      SELECT COALESCE(sum(i.quantidade),0) FROM producao.apontamento_item i
      JOIN producao.produto p ON p.id=i.produto_id
      WHERE i.apontamento_dia_id=:id AND lower(trim(p.familia))<>lower(:non_post)
        """), {"id": marker["id"], "non_post": NON_POST_FAMILY}).scalar_one()
        people = _labor(db, unit, dia)["homens"]
        try:
            assert_can_confirm(int(posts), people, step=unit["passo_homens"])
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from exc
        db.execute(text("""
          UPDATE producao.apontamento_dia SET status='confirmado',confirmado_em=now(),confirmado_por=:user
          WHERE id=:id
        """), {"user": user.username, "id": marker["id"]})
        db.commit()
    except Exception:
        db.rollback()
        raise
    return _production(db, unit, dia)


@router.get("/mensal")
def monthly(unidade: str, mes: str = Query(pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
            db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    unit = _unit(db, unidade)
    _authorize(db, unit, user)
    dates = [row[0] for row in db.execute(text("""
      SELECT d.data FROM producao.apontamento_dia d
      JOIN producao.apontamento_item i ON i.apontamento_dia_id=d.id
      JOIN producao.produto p ON p.id=i.produto_id
      WHERE d.unidade_id=:uid AND to_char(d.data,'YYYY-MM')=:mes AND d.data<=:today
        AND lower(trim(p.familia))<>lower(:non_post)
      GROUP BY d.data HAVING sum(i.quantidade)>0 ORDER BY d.data
    """), {"uid": unit["id"], "mes": mes,
           "non_post": NON_POST_FAMILY,
           "today": datetime.now(ZoneInfo("America/Sao_Paulo")).date()})]
    days = []
    for day in dates:
        production = _production(db, unit, day)
        labor = _labor(db, unit, day)
        days.append({"data": day, "status": production["status"], "postes": production["postes"],
                     "homens": labor["homens"], "mod": labor["mod"], "moi": labor["moi"],
                     "postes_homem": posts_per_person(production["postes"], labor["homens"], step=unit["passo_homens"]),
                     "volume_m3": production["volume_m3"],
                     "produtos_sem_volume": production["produtos_sem_volume"],
                     "areas": labor["areas"],
                     "itens": production["itens"]})
    meta = db.execute(text("""
      SELECT dias_uteis,pcs_dia,meta_postes_homem FROM producao.meta_producao
      WHERE unidade_id=:uid AND ano_mes=:mes
    """), {"uid": unit["id"], "mes": mes}).mappings().first()
    total_posts = sum(day["postes"] for day in days)
    total_people = sum((day["homens"] for day in days), Decimal("0"))
    total_volume = (sum((day["volume_m3"] for day in days), Decimal("0"))
                    if all(day["volume_m3"] is not None for day in days) else None)
    return {"unidade": unit, "mes": mes, "meta": dict(meta) if meta else None,
            "meta_total": monthly_target(meta["dias_uteis"], meta["pcs_dia"]) if meta else None,
            "dias": days, "dias_producao": len(days), "postes": total_posts,
            "media_postes_dia": round(total_posts / len(days), 2) if days else None,
            "homens": total_people, "postes_homem": posts_per_person(total_posts, total_people, step=unit["passo_homens"]),
            "volume_m3": total_volume,
            "produtos_sem_volume": sum(day["produtos_sem_volume"] for day in days),
            "pendencias": [day["data"] for day in days if day["postes"] and not day["homens"]]}


@router.get("/exportar/lt-pro-04.xlsx")
def export_xlsx(unidade: str, mes: str = Query(pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
                db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    from app.services.pederneiras_exports import xlsx
    data = monthly(unidade, mes, db, user)
    stream = BytesIO(xlsx(data, _product_rows(db, data["unidade"])))
    return StreamingResponse(stream,
                             media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="LT-PRO-04-{unidade}-{mes}.xlsx"'})


@router.get("/exportar/lt-pro-04.pdf")
def export_pdf(unidade: str, mes: str = Query(pattern=r"^\d{4}-(0[1-9]|1[0-2])$"),
               db: Session = Depends(get_db), user: Usuario = Depends(get_current_user)):
    from app.services.pederneiras_exports import pdf
    data = monthly(unidade, mes, db, user)
    stream = BytesIO(pdf(data, _product_rows(db, data["unidade"])))
    return StreamingResponse(stream, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="LT-PRO-04-{unidade}-{mes}.pdf"'})


