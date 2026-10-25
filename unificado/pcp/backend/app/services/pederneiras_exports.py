"""LT-PRO-04 exports from the same monthly payload used by the UI."""
from __future__ import annotations

from datetime import date
from io import BytesIO


def _grid(month: dict, products: list[dict]):
    days = month["dias"]
    quantity = {(item["produto_id"], day["data"]): item["quantidade"]
                for day in days for item in day["itens"]}
    rows = []
    for product in products:
        values = [quantity.get((product["id"], day["data"]), 0) for day in days]
        if any(values):
            rows.append((product, values, sum(values)))
    return days, rows


def xlsx(month: dict, products: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    days, rows = _grid(month, products)
    book = Workbook()
    sheet = book.active
    sheet.title = "LT-PRO-04"
    sheet.append([f"LT-PRO-04 · {month['unidade']['nome']} · {month['mes']}"])
    sheet.append(["Produto", "Cód. Pro", "Cód. SMA",
                  *[day["data"].strftime("%d/%m") for day in days], "Total"])
    for product, values, total in rows:
        sheet.append([product["descricao"], product["cod_pro"], product["cod_sma"], *values, total])
    footer = [
        ("Total postes", [day["postes"] for day in days], month["postes"]),
        ("Homens", [day["homens"] for day in days], month["homens"]),
        ("MOD", [day["mod"] for day in days], sum((day["mod"] for day in days), 0)),
        ("MOI", [day["moi"] for day in days], sum((day["moi"] for day in days), 0)),
        ("Postes/homem", [day["postes_homem"] for day in days], None),
        ("m³", [day["volume_m3"] for day in days], None),
    ]
    for label, values, total in footer:
        sheet.append([label, None, None, *values, total])
    sheet.freeze_panes = "D3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(4 + len(days))}{2 + len(rows)}"
    sheet.column_dimensions["A"].width = 49
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 16
    for column in range(4, 5 + len(days)):
        sheet.column_dimensions[get_column_letter(column)].width = 11
    for cell in sheet[2]:
        cell.fill = PatternFill("solid", fgColor="17324F")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for index, day in enumerate(days, 4):
        if day["data"].weekday() >= 5:
            for row in sheet.iter_rows(min_col=index, max_col=index, min_row=2):
                row[0].fill = PatternFill("solid", fgColor="FFF0DA")
    for row in sheet.iter_rows(min_row=3 + len(rows)):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EDF3F8")
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()


def pdf(month: dict, products: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Spacer, Paragraph, Table, TableStyle

    days, rows = _grid(month, products)
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=landscape(A3), leftMargin=24,
                                 rightMargin=24, topMargin=25, bottomMargin=25)
    styles = getSampleStyleSheet()
    story = [Paragraph(f"LT-PRO-04 · Controle de produção · {month['unidade']['nome']} · {month['mes']}", styles["Title"]),
             Paragraph("Somente dias com produção até a data atual. Volume ausente é exibido como -.", styles["Normal"]),
             Spacer(1, 12)]
    table = [["Produto", "Cód. Pro", "Cód. SMA", *[day["data"].strftime("%d/%m") for day in days], "Total"]]
    for product, values, total in rows:
        table.append([product["descricao"], product["cod_pro"] or "", product["cod_sma"] or "",
                      *[value or "" for value in values], total])
    for label, field in (("Total postes", "postes"), ("Homens", "homens"),
                         ("MOD", "mod"), ("MOI", "moi"), ("Postes/homem", "postes_homem"),
                         ("m³", "volume_m3")):
        values = [day[field] if day[field] is not None else "-" for day in days]
        total = month["postes"] if field == "postes" else month["homens"] if field == "homens" else ""
        table.append([label, "", "", *values, total])
    page_width = landscape(A3)[0] - 48
    day_width = min(43, max(23, (page_width - 300) / (len(days) + 1)))
    description_width = page_width - 2 * 58 - (len(days) + 1) * day_width
    widths = [description_width, 58, 58, *([day_width] * (len(days) + 1))]
    report = Table(table, colWidths=widths, repeatRows=1, hAlign="LEFT")
    footer_start = 1 + len(rows)
    rules = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17324F")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, footer_start), (-1, -1), "Helvetica-Bold"),
        ("BACKGROUND", (0, footer_start), (-1, -1), colors.HexColor("#EDF3F8")),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#D9E2EA")),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("LEADING", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (3, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]
    for index, day in enumerate(days, 3):
        if day["data"].weekday() >= 5:
            rules.append(("BACKGROUND", (index, 1), (index, footer_start - 1), colors.HexColor("#FFF0DA")))
    report.setStyle(TableStyle(rules))
    story.append(report)
    document.build(story)
    return buffer.getvalue()
