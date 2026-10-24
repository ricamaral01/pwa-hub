"""Product catalogue import parsing and deterministic preview."""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation


FIELDS = ("familia", "letra", "descricao", "cod_pro", "cod_sma", "volume_m3", "ativo", "pendencia")


def _key(value: object) -> str:
    plain = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", plain)


HEADERS = {_key(field): field for field in FIELDS}
HEADERS.update({"codproducao": "cod_pro", "codigosma": "cod_sma", "m3peca": "volume_m3",
                "situacao": "pendencia", "produto": "descricao", "familia": "familia"})


def _decimal(value: object) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    raw = str(value).strip().replace(" ", "")
    if "," in raw:
        raw = raw.replace(".", "").replace(",", ".")
    try:
        number = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"Volume inválido: {value}") from exc
    if not number.is_finite() or not 0 <= number <= Decimal("9999.9999"):
        raise ValueError(f"Volume fora da faixa: {value}")
    return number


def _active(value: object) -> bool:
    norm = _key(value)
    if norm in {"", "1", "true", "sim", "s", "ativo"}:
        return True
    if norm in {"0", "false", "nao", "n", "inativo"}:
        return False
    raise ValueError(f"Ativo inválido: {value}")


def _clean_row(source: dict, line: int) -> dict:
    mapped = {HEADERS[_key(key)]: value for key, value in source.items() if _key(key) in HEADERS}
    row = {field: str(mapped.get(field) or "").strip() or None for field in FIELDS}
    if not row["familia"] or not row["descricao"]:
        raise ValueError(f"Linha {line}: família e descrição obrigatórias")
    row["volume_m3"] = _decimal(mapped.get("volume_m3"))
    row["ativo"] = _active(mapped.get("ativo"))
    return row


def parse_catalog(name: str, content: bytes) -> list[dict]:
    if len(content) > 5_000_000:
        raise ValueError("Arquivo excede 5 MB")
    lower = name.lower()
    if lower.endswith(".csv"):
        stream = io.StringIO(content.decode("utf-8-sig"))
        sample = stream.read(4096)
        stream.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,")
            reader = csv.DictReader(stream, dialect=dialect)
        except csv.Error:
            reader = csv.DictReader(stream, delimiter=";")
        raw = list(reader)
    elif lower.endswith(".xlsx"):
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        header = next(values, None)
        if not header:
            raise ValueError("Planilha vazia")
        raw = [dict(zip(header, row)) for row in values if any(value is not None for value in row)]
    else:
        raise ValueError("Formato aceito: CSV ou XLSX")
    if not raw:
        raise ValueError("Arquivo sem produtos")
    rows = [_clean_row(item, i) for i, item in enumerate(raw, 2)]
    codes = [row["cod_pro"] for row in rows if row["cod_pro"]]
    if len(codes) != len(set(codes)):
        raise ValueError("Código Pro repetido no arquivo")
    return rows


def preview(rows: list[dict], existing: list[dict]) -> list[dict]:
    by_code = {str(row["cod_pro"]): row for row in existing if row["cod_pro"]}
    by_description = {}
    for row in existing:
        key = (_key(row["familia"]), _key(row["descricao"]))
        by_description.setdefault(key, []).append(row)
    result = []
    for index, row in enumerate(rows, 2):
        found = by_code.get(row["cod_pro"]) if row["cod_pro"] else None
        candidates = by_description.get((_key(row["familia"]), _key(row["descricao"])), [])
        if found is None and len(candidates) == 1:
            found = candidates[0]
        if found is None and len(candidates) > 1:
            result.append({"linha": index, "acao": "conflito", "motivo": "Descrição ambígua", "produto": row})
            continue
        changes = {field: {"de": found[field], "para": row[field]}
                   for field in FIELDS if found is not None and found[field] != row[field]}
        result.append({"linha": index, "acao": "criar" if found is None else "atualizar" if changes else "igual",
                       "produto_id": found["id"] if found else None, "produto": row, "alteracoes": changes})
    return result
